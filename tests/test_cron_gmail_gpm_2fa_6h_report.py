import json
import logging
import sqlite3
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "deploy" / "hermes-home" / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import cron_gmail_gpm_2fa_6h_report as report_mod

LOGGER_NAME = "cron_gmail_gpm_2fa_6h_report"
LONG_SECRET = "SEC12345678901234567890123456789012"


@pytest.fixture
def paths(tmp_path):
    """Trỏ toàn bộ nguồn dữ liệu vào tmp_path (mặc định: không file nào tồn tại)."""
    p = {
        "STATE_FILE": tmp_path / "state.json",
        "GPM_DB_PATH": tmp_path / "profile.db",
        "MASTER_EXCEL": tmp_path / "master.xlsx",
        "CLEAN_V2_EXCEL": tmp_path / "clean.xlsx",
        "NURTURE_STATE_FILE": tmp_path / "nurture.json",
    }
    with patch.multiple(report_mod, **p):
        yield p


def write_state(path, events):
    path.write_text(json.dumps({"events": events}), encoding="utf-8")


def make_db(path, rows):
    with sqlite3.connect(path) as c:
        c.execute("CREATE TABLE Profiles(Id,Name,ProfilePath,GroupId)")
        c.executemany("INSERT INTO Profiles VALUES(?,?,?,?)", rows)


def make_master(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_All"
    ws.append(["STT", "Email", "Password", "Recovery", "2FA_Secret"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def make_clean(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["STT", "Email", "Password", "2FA_Secret"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def metrics(caplog):
    out = []
    for rec in caplog.records:
        msg = rec.getMessage()
        if msg.startswith("[TELEMETRY_METRIC] "):
            out.append(json.loads(msg[len("[TELEMETRY_METRIC] "):]))
    return out


def recent_iso(hours=1):
    return (datetime.now() - timedelta(hours=hours)).isoformat()


# ---------- main(): luồng end-to-end ----------


def test_report_rendering_with_events(capsys, paths):
    write_state(paths["STATE_FILE"], [
        {"email": "test1@gmail.com", "machine": 18, "status": "SUCCESS",
         "secret_key": "ABCDEF1234567890", "at": recent_iso(1)},
        {"email": "test2@gmail.com", "machine": 26, "status": "ALREADY_ACTIVE", "at": recent_iso(1)},
        {"email": "old@gmail.com", "machine": 10, "status": "SUCCESS", "at": recent_iso(8)},
    ])

    assert report_mod.main() == 0

    out = capsys.readouterr().out
    assert "📊 [BÁO CÁO 6H: 2FA GMAIL QUA GPM PROFILE]" in out
    assert "⚡ Tiến độ xử lý trong 6h vừa qua: 2 sự kiện" in out
    assert "M18 | test1@gmail.com: SUCCESS (ABCDEF12...7890)" in out
    assert "M26 | test2@gmail.com: ALREADY_ACTIVE" in out
    assert "old@gmail.com" not in out


def test_silent_watchdog_when_zero_events(capsys, paths, caplog):
    write_state(paths["STATE_FILE"], [{"email": "old@gmail.com", "machine": 10, "status": "SUCCESS",
                                       "at": recent_iso(8)}])
    with caplog.at_level(logging.INFO, logger=LOGGER_NAME):
        assert report_mod.main() == 0

    assert capsys.readouterr().out == ""
    (metric,) = metrics(caplog)
    assert metric["silent"] is True
    assert metric["recent_events"] == 0


def test_silent_when_state_file_missing(capsys, paths):
    assert report_mod.main() == 0
    assert capsys.readouterr().out == ""


def test_full_pipeline_counts_and_telemetry(capsys, paths, caplog):
    make_db(paths["GPM_DB_PATH"], [
        (1, "has2fa@gmail.com", "p", 10),
        (2, "clean2fa@gmail.com", "p", 10),
        (3, "ready@gmail.com", "p", 10),
        (4, "soak@gmail.com", "p", 10),
        (5, "khoale123@gmail.com", "p", 10),
        (6, "no-email-profile", "p", 10),
        (7, "other-group@gmail.com", "p", 11),
    ])
    make_master(paths["MASTER_EXCEL"], [["1", "Has2FA@Gmail.com", "", "", LONG_SECRET]])
    make_clean(paths["CLEAN_V2_EXCEL"], [["1", "clean2fa@gmail.com", "", "CLEANSECRET"]])
    paths["NURTURE_STATE_FILE"].write_text(json.dumps({
        "ready@gmail.com": {"last_nurtured": "2026-10-01", "status": "success"},
        "soak@gmail.com": {"last_nurtured": "2026-10-01", "status": "failed"},
    }), encoding="utf-8")
    write_state(paths["STATE_FILE"], [{"email": "x@gmail.com", "machine": 3, "status": "FAILED",
                                       "details": "d" * 100, "at": recent_iso(2)}])

    with caplog.at_level(logging.INFO, logger=LOGGER_NAME):
        assert report_mod.main() == 0

    out = capsys.readouterr().out
    assert "Tổng số profile GPM (Group 10): 6" in out
    assert "Đã bật 2FA thành công: 2/6 (33.3%)" in out
    assert "Đang ngâm (chờ phiên nuôi): 1" in out
    assert "Sẵn sàng bật 2FA: 1" in out
    assert "M03 | x@gmail.com: FAILED - " + "d" * 60 in out
    assert "d" * 61 not in out
    assert "⚠️" not in out

    (metric,) = metrics(caplog)
    assert metric["event"] == "gmail_2fa_6h_report"
    assert (metric["total"], metric["already_2fa"], metric["soaking"], metric["ready"], metric["skipped"]) == (6, 2, 1, 1, 2)
    assert metric["silent"] is False
    assert metric["failed_sources"] == []


def test_log_goes_to_stderr_not_stdout(capsys, paths):
    write_state(paths["STATE_FILE"], [{"email": "a@gmail.com", "machine": 1, "status": "SUCCESS", "at": recent_iso()}])
    report_mod.logger.handlers.clear()
    report_mod.main()
    captured = capsys.readouterr()
    assert "[TELEMETRY_METRIC]" in captured.err
    assert "[TELEMETRY_METRIC]" not in captured.out


# ---------- xử lý lỗi: log rõ ràng, không crash, đánh dấu nguồn hỏng ----------


def test_corrupted_state_and_sqlite_error_are_logged_and_silent(capsys, paths, caplog):
    paths["STATE_FILE"].write_text("{not valid json", encoding="utf-8")
    paths["GPM_DB_PATH"].write_text("corrupt sqlite header", encoding="utf-8")

    with caplog.at_level(logging.INFO, logger=LOGGER_NAME):
        assert report_mod.main() == 0

    assert capsys.readouterr().out == ""
    errors = [r.getMessage() for r in caplog.records if r.levelno >= logging.ERROR]
    assert any("STATE_FILE" in m for m in errors)
    assert any("GPM_DB" in m for m in errors)
    (metric,) = metrics(caplog)
    assert set(metric["failed_sources"]) == {"STATE_FILE", "GPM_DB"}


def test_degraded_sources_are_flagged_in_report(capsys, paths):
    paths["GPM_DB_PATH"].write_text("corrupt sqlite header", encoding="utf-8")
    paths["MASTER_EXCEL"].write_text("not a zip", encoding="utf-8")
    paths["NURTURE_STATE_FILE"].write_text("[1, 2]", encoding="utf-8")
    write_state(paths["STATE_FILE"], [{"email": "a@gmail.com", "machine": 1, "status": "SUCCESS", "at": recent_iso()}])

    assert report_mod.main() == 0

    out = capsys.readouterr().out
    assert "⚠️ Nguồn dữ liệu lỗi" in out
    for name in ("MASTER_EXCEL", "NURTURE_STATE_FILE", "GPM_DB"):
        assert name in out
    assert "Tổng số profile GPM (Group 10): 0" in out
    assert "(0%)" in out


def test_non_dict_state_file_is_rejected(paths, caplog):
    paths["STATE_FILE"].write_text("[]", encoding="utf-8")
    failed = []
    with caplog.at_level(logging.ERROR, logger=LOGGER_NAME):
        assert report_mod.load_json_file(paths["STATE_FILE"], "STATE_FILE", failed) == {}
    assert failed == ["STATE_FILE"]
    assert "không phải JSON object" in caplog.text


def test_unreadable_json_path_is_reported(tmp_path, caplog):
    failed = []
    with caplog.at_level(logging.ERROR, logger=LOGGER_NAME):
        assert report_mod.load_json_file(tmp_path, "STATE_FILE", failed) == {}  # directory -> OSError
    assert failed == ["STATE_FILE"]


def test_missing_file_is_not_a_failure(tmp_path):
    failed = []
    assert report_mod.load_json_file(tmp_path / "nope.json", "STATE_FILE", failed) == {}
    assert failed == []


def test_master_excel_missing_sheet_is_warning_not_failure(paths, caplog):
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    wb.save(paths["MASTER_EXCEL"])
    failed = []
    with caplog.at_level(logging.WARNING, logger=LOGGER_NAME):
        assert report_mod.load_master_2fa(failed) == {}
    assert failed == []
    assert "thiếu sheet Master_All" in caplog.text


def test_excel_short_rows_and_blank_secrets_are_ignored(paths):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_All"
    ws.append(["STT", "Email", "Password", "Recovery", "2FA_Secret"])
    ws.append(["1", "short@gmail.com"])
    ws.append(["2", "blank@gmail.com", "", "", None])
    ws.append(["3", None, "", "", "ORPHANSECRET"])
    ws.append(["4", "  Good@Gmail.com ", "", "", " SECRET "])
    wb.save(paths["MASTER_EXCEL"])
    failed = []
    assert report_mod.load_master_2fa(failed) == {"good@gmail.com": "SECRET"}
    assert failed == []


# ---------- hàm đơn vị ----------


@pytest.mark.parametrize(
    "name,expected",
    [
        ("Real.User+1@Gmail.com", "real.user+1@gmail.com"),
        ("prefix 12 abc@gmail.com suffix", "abc@gmail.com"),
        ("someone@hotmail.com", None),
        ("", None),
        (None, None),
    ],
)
def test_extract_email(name, expected):
    assert report_mod.extract_email(name) == expected


@pytest.mark.parametrize("secret", ["None", "null", ""])
def test_classify_treats_placeholder_secret_as_missing(secret):
    stats = report_mod.classify_profiles(["a@gmail.com"], {"a@gmail.com": secret}, {}, {})
    assert stats["already_2fa"] == 0
    assert stats["soaking"] == 1


def test_classify_master_takes_precedence_and_handles_bad_nurture_entry():
    names = ["a@gmail.com", "b@gmail.com", "c@gmail.com"]
    stats = report_mod.classify_profiles(
        names,
        {"a@gmail.com": "S"},
        {"a@gmail.com": "other"},
        {"b@gmail.com": "garbage-not-a-dict", "c@gmail.com": {"last_nurtured": "x", "status": "success"}},
    )
    assert (stats["already_2fa"], stats["soaking"], stats["ready"]) == (1, 1, 1)


def test_classify_empty_input():
    assert report_mod.classify_profiles([], {}, {}, {}) == {
        "total": 0, "already_2fa": 0, "soaking": 0, "ready": 0, "skipped": 0,
    }


def test_filter_recent_events_edge_cases(caplog):
    now = datetime.now()
    events = [
        {"email": "ok", "at": (now - timedelta(minutes=5)).isoformat()},
        {"email": "z-suffix", "at": (now - timedelta(minutes=5)).isoformat() + "Z"},
        {"email": "aware", "at": datetime.now(timezone.utc).isoformat()},
        {"email": "old", "at": (now - timedelta(hours=7)).isoformat()},
        {"email": "no-at"},
        {"email": "bad-at", "at": "not-a-date"},
        "not-a-dict",
    ]
    with caplog.at_level(logging.WARNING, logger=LOGGER_NAME):
        result = report_mod.filter_recent_events(events, now - timedelta(hours=6))
    assert [e["email"] for e in result] == ["ok", "z-suffix", "aware"]
    assert "Lỗi parse timestamp" in caplog.text
    assert "Bỏ qua event không phải dict" in caplog.text


def test_filter_recent_events_rejects_non_list(caplog):
    with caplog.at_level(logging.ERROR, logger=LOGGER_NAME):
        assert report_mod.filter_recent_events({"a": 1}, datetime.now()) == []
    assert "không phải list" in caplog.text


@pytest.mark.parametrize(
    "event,expected",
    [
        ({"email": "a@x", "machine": 5, "status": "SUCCESS", "secret_key": "SHORT"}, "M05 | a@x: SUCCESS (SHORT)"),
        ({"email": "a@x", "machine": 5, "status": "SUCCESS"}, "M05 | a@x: SUCCESS"),
        ({"email": "a@x", "machine": 5, "status": "SUCCESS", "secret_key": None}, "M05 | a@x: SUCCESS"),
        ({"email": "a@x", "status": "ALREADY_ACTIVE"}, "M?? | a@x: ALREADY_ACTIVE"),
        ({"email": "a@x", "machine": "abc", "status": "FAILED", "details": None}, "M?? | a@x: FAILED - "),
        ({"email": "a@x", "machine": "12", "status": "FAILED", "details": "boom"}, "M12 | a@x: FAILED - boom"),
    ],
)
def test_format_event(event, expected):
    assert expected in report_mod.format_event(event)


def test_render_report_zero_total_has_no_division_error():
    stats = {"total": 0, "already_2fa": 0, "soaking": 0, "ready": 0, "skipped": 0}
    out = report_mod.render_report(stats, [], [])
    assert "0/0 (0%)" in out
    assert "0 sự kiện" in out


def test_emit_metric_is_valid_json_with_unicode(caplog):
    with caplog.at_level(logging.INFO, logger=LOGGER_NAME):
        report_mod.emit_metric("evt", note="tiếng Việt", n=1)
    (metric,) = metrics(caplog)
    assert metric == {"event": "evt", "note": "tiếng Việt", "n": 1}
    assert "tiếng Việt" in caplog.text


def test_setup_logging_is_idempotent():
    report_mod.logger.handlers.clear()
    report_mod.setup_logging()
    report_mod.setup_logging()
    assert len(report_mod.logger.handlers) == 1
    report_mod.logger.handlers.clear()
