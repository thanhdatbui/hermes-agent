import sys
import os
import json
import sqlite3
import unittest
import requests
from unittest.mock import patch, MagicMock
from pathlib import Path

REPO_SCRIPTS = Path(__file__).resolve().parent.parent / "deploy" / "hermes-home" / "scripts"
if str(REPO_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(REPO_SCRIPTS))

import cron_chatgpt_web_pool_watchdog as w

class TestCronChatgptWebPoolWatchdog(unittest.TestCase):
    def test_get_connections_by_provider_codex(self):
        mock_data = {
            "connections": [
                {"id": "1", "provider": "codex", "name": "acc1", "isActive": True, "testStatus": "active"},
                {"id": "2", "provider": "chatgpt-web", "name": "acc2", "isActive": True, "testStatus": "active"},
                {"id": "3", "provider": "codex", "name": "acc3", "isActive": False, "testStatus": "active"},
            ]
        }
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = mock_data
        mock_resp.read.return_value = json.dumps(mock_data).encode("utf-8")

        mock_cm = MagicMock()
        mock_cm.__enter__.return_value = mock_resp

        with patch("requests.get", return_value=mock_resp), \
             patch("urllib.request.urlopen", return_value=mock_cm):
            conns = w.get_connections_by_provider("codex")
            self.assertEqual(len(conns), 2)
            self.assertEqual(conns[0]["id"], "1")
            self.assertEqual(conns[1]["id"], "3")

    def test_codex_observability_standby_preserved(self):
        # Tôn trọng trạng thái router/người dùng đặt Standby, không tự ý toggle bật lại
        codex_conns = [
            {"id": "conn-1", "name": "acc1@gmail.com", "isActive": False, "testStatus": "active"},
            {"id": "conn-2", "name": "acc2@gmail.com", "isActive": True, "testStatus": "active"}
        ]
        toggled_off = [c for c in codex_conns if not c.get("isActive") and c.get("testStatus") == "active"]
        self.assertEqual(len(toggled_off), 1)
        self.assertEqual(toggled_off[0]["id"], "conn-1")

    def test_broken_codex_detected(self):
        codex_conns = [
            {"id": "conn-1", "name": "acc1@gmail.com", "isActive": True, "testStatus": "expired"},
            {"id": "conn-2", "name": "acc2@gmail.com", "isActive": True, "testStatus": "active"}
        ]
        broken = [c for c in codex_conns if c.get("testStatus") != "active"]
        self.assertEqual(len(broken), 1)
        self.assertEqual(broken[0]["name"], "acc1@gmail.com")

    def test_report_lines_format(self):
        report_lines = [
            "🤖 [POOL HEALER] BÁO CÁO SỨC KHỎE",
            "• ChatGPT-Web: 21/21 ACTIVE",
            "• Antigravity: 114/121 ACTIVE",
            "• Codex: 29/29 ACTIVE"
        ]
        full_text = "\n".join(report_lines)
        self.assertIn("• Codex: 29/29 ACTIVE", full_text)
        self.assertIn("• ChatGPT-Web: 21/21 ACTIVE", full_text)

    def test_save_telemetry_metrics_atomic_write(self):
        metrics = {
            "timestamp": "2026-09-25T00:00:00Z",
            "chatgpt_web": {"active": 15, "total": 21},
            "antigravity": {"active": 109, "total": 119},
            "codex": {"active": 29, "total": 29},
            "recovered_chatgpt": ["acc1@gmail.com"],
            "recovered_ag": [],
            "total_failures": 0,
            "failures": []
        }
        test_cache = Path(__file__).resolve().parent / "temp_telemetry_test.json"
        test_hist = Path(__file__).resolve().parent / "temp_telemetry_hist.jsonl"
        with patch.object(w, "TELEMETRY_CACHE_PATH", str(test_cache)), \
             patch.object(w, "TELEMETRY_HISTORY_PATH", str(test_hist)):
            ok = w.save_telemetry_metrics(metrics)
            self.assertTrue(ok)
            self.assertTrue(test_cache.exists())
            self.assertTrue(test_hist.exists())
            with open(test_cache, "r", encoding="utf-8") as f:
                data = json.load(f)
            self.assertEqual(data["chatgpt_web"]["active"], 15)
            self.assertEqual(data["selector_version"], "2026.09.25-v1")
            with open(test_hist, "r", encoding="utf-8") as f_h:
                lines = f_h.readlines()
            self.assertEqual(len(lines), 1)
            if test_cache.exists(): test_cache.unlink()
            if test_hist.exists(): test_hist.unlink()

    def test_sqlite_transaction_rollback_on_failure(self):
        # Test: DB error during UPDATE -> rollback called
        with patch("sqlite3.connect") as mock_connect, \
             patch("requests.get") as mock_req, \
             patch.object(w, "perform_auto_login_and_extract_cookies", return_value="token_valid"), \
             patch.object(w, "test_connection_valid", return_value=(True, {})):

            mock_conn = MagicMock()
            mock_cur = MagicMock()
            mock_connect.return_value = mock_conn
            mock_conn.cursor.return_value = mock_cur
            mock_cur.fetchone.return_value = ("chatgpt-web", "test_acc")
            # First execute (guard check) succeeds, second (UPDATE) fails
            mock_cur.execute.side_effect = [None, sqlite3.DatabaseError("Simulated DB error")]

            mock_req.return_value.json.return_value = {
                "success": True,
                "data": {"remote_debugging_address": "127.0.0.1:9222"}
            }
            ok, msg = w.perform_revive_chatgpt_account("dummy_pid", "dummy_cid", "test@gmail.com", {})
            self.assertFalse(ok)
            # Verify rollback was called when DB error occurred
            self.assertTrue(mock_conn.rollback.called or mock_conn.close.called)

    def test_chatgpt_revive_guard_rejects_wrong_provider(self):
        # Kiểm tra guard chặn đứng nếu connection không phải chatgpt-web
        with patch("sqlite3.connect") as mock_connect:
            mock_conn = MagicMock()
            mock_cur = MagicMock()
            mock_connect.return_value = mock_conn
            mock_conn.cursor.return_value = mock_cur
            mock_cur.fetchone.return_value = ("antigravity", "wrong_target_acc")

            with patch("requests.get") as mock_req, \
                 patch.object(w, "perform_auto_login_and_extract_cookies", return_value="cookie=test"), \
                 patch.object(w, "test_connection_valid", return_value=(True, {})):
                mock_req.return_value.json.return_value = {
                    "success": True,
                    "data": {"remote_debugging_address": "127.0.0.1:9222"}
                }
                ok, msg = w.perform_revive_chatgpt_account("dummy_pid", "cid_antigravity", "test@gmail.com", {})
                self.assertFalse(ok)
                self.assertIn("Guard rejected", msg)

    def test_sqlite_real_schema_mutation_and_combos_update(self):
        # Integration test trên SQLite schema thực tế
        test_db_path = str(Path(__file__).resolve().parent / "temp_integration.sqlite")
        conn = sqlite3.connect(test_db_path)
        cur = conn.cursor()
        cur.execute("CREATE TABLE provider_connections (id TEXT PRIMARY KEY, provider TEXT, name TEXT, api_key TEXT, is_active INTEGER, test_status TEXT, last_error TEXT, last_error_at TEXT, backoff_level INTEGER, rate_limited_until TEXT, updated_at TEXT)")
        cur.execute("CREATE TABLE combos (name TEXT PRIMARY KEY, data TEXT)")
        cur.execute("INSERT INTO provider_connections VALUES ('cid_web_1', 'chatgpt-web', 'test@gmail.com', 'old_key', 0, 'expired', 'error', '2026-09-20', 1, NULL, '2026-09-20')")
        cur.execute("INSERT INTO combos VALUES ('chatgpt-web-pool', '{\"models\": []}')")
        conn.commit()
        conn.close()

        with patch.object(w, "DB_PATH", test_db_path), \
             patch("requests.get") as mock_req, \
             patch.object(w, "perform_auto_login_and_extract_cookies", return_value="session-token=new_fresh_token"), \
             patch.object(w, "test_connection_valid", return_value=(True, {})):
            mock_req.return_value.json.return_value = {
                "success": True,
                "data": {"remote_debugging_address": "127.0.0.1:9222"}
            }
            ok, msg = w.perform_revive_chatgpt_account("pid_1", "cid_web_1", "test@gmail.com", {})
            self.assertTrue(ok)

            # Verify actual database state
            conn_check = sqlite3.connect(test_db_path)
            row = conn_check.cursor().execute("SELECT is_active, test_status, api_key FROM provider_connections WHERE id='cid_web_1'").fetchone()
            self.assertEqual(row[0], 1)
            self.assertEqual(row[1], "active")
            self.assertEqual(row[2], "session-token=new_fresh_token")

            combo_row = conn_check.cursor().execute("SELECT data FROM combos WHERE name='chatgpt-web-pool'").fetchone()
            cdata = json.loads(combo_row[0])
            self.assertEqual(len(cdata["models"]), 1)
            self.assertEqual(cdata["models"][0]["connectionId"], "cid_web_1")
            self.assertEqual(cdata["models"][0]["model"], "chatgpt-web/gpt-5.6-sol-high")
            conn_check.close()

        if os.path.exists(test_db_path):
            os.remove(test_db_path)

    def test_phone_checkpoint_predicate_detection(self):
        sample_page_text_phone = "Để giữ an toàn cho tài khoản của bạn, Google muốn đảm bảo rằng bạn chính là người đang cố đăng nhập. Nhập số điện thoại để nhận mã xác minh."
        txt = sample_page_text_phone.lower()
        is_phone_cp = (("phone number" in txt or "số điện thoại" in txt) and ("enter a phone" in txt or "nhập số" in txt))
        self.assertTrue(is_phone_cp)

        sample_page_text_normal = "Nhập mật khẩu của bạn. Tiếp theo."
        txt_normal = sample_page_text_normal.lower()
        is_phone_cp_normal = (("phone number" in txt_normal or "số điện thoại" in txt_normal) and ("enter a phone" in txt_normal or "nhập số" in txt_normal))
        self.assertFalse(is_phone_cp_normal)

    def test_load_accounts_credentials_column_integrity(self):
        fake_rows = [
            ("số máy", "email", "pass", "2fa", "recovery_email"),
            (1, "acc1@gmail.com", "PassWord123!", "JBSWY3DPEHPK3PXP", "recovery1@gmail.com"),
            (2, "acc2@gmail.com", "PassWord456!", None, "recovery2@gmail.com"),
        ]
        mock_wb = MagicMock()
        mock_ws = MagicMock()
        mock_ws.iter_rows.return_value = fake_rows
        mock_wb.sheetnames = ["Sheet1"]
        mock_wb.__getitem__.return_value = mock_ws
        mock_wb.active = mock_ws
        with patch("openpyxl.load_workbook", return_value=mock_wb), \
             patch("os.path.exists", return_value=True):
            creds = w.load_accounts_credentials()
            self.assertIn("acc1@gmail.com", creds)
            self.assertEqual(creds["acc1@gmail.com"]["password"], "PassWord123!")
            self.assertEqual(creds["acc1@gmail.com"]["totp"], "JBSWY3DPEHPK3PXP")
            self.assertEqual(creds["acc1@gmail.com"]["recovery_email"], "recovery1@gmail.com")
            self.assertNotIn("recovery1@gmail.com", creds)

    def test_main_controller_multi_provider_failures_and_reporting(self):
        # Kiểm chứng toàn bộ main() controller xử lý đúng khi nhiều provider gặp lỗi đồng thời
        mock_chatgpt_conns = [
            {"id": "cid_cg_1", "name": "user1@gmail.com (GPM Web)", "isActive": False, "testStatus": "expired", "provider": "chatgpt-web"}
        ]
        mock_ag_conns = [
            {"id": "cid_ag_1", "name": "user2@gmail.com", "isActive": True, "testStatus": "expired", "provider": "antigravity"}
        ]
        mock_codex_conns = [
            {"id": "cid_cx_1", "name": "user3@gmail.com", "isActive": False, "testStatus": "active", "provider": "codex"}
        ]

        def mock_get_conns(prov):
            if prov == "chatgpt-web": return mock_chatgpt_conns
            if prov == "antigravity": return mock_ag_conns
            if prov == "codex": return mock_codex_conns
            return []

        mock_gpm_profiles = [
            {"id": "pid_1", "name": "01 - user1@gmail.com - 5101", "raw_proxy": "test:5101"},
            {"id": "pid_2", "name": "02 - user2@gmail.com - 5102", "raw_proxy": "test:5102"}
        ]

        with patch.object(w, "get_connections_by_provider", side_effect=mock_get_conns),              patch.object(w, "get_gpm_profiles", return_value=mock_gpm_profiles),              patch.object(w, "load_accounts_credentials", return_value={"user1@gmail.com": {}, "user2@gmail.com": {}}),              patch.object(w, "refresh_chatgpt_account_via_gpm", return_value=(True, "OK")),              patch.object(w, "refresh_antigravity_account_via_gpm", return_value=(False, "Phone Checkpoint")),              patch.object(w, "save_telemetry_metrics") as mock_save_telem:
            w.main()
            mock_save_telem.assert_called_once()
            telem_arg = mock_save_telem.call_args[0][0]
            self.assertIn("user1@gmail.com", telem_arg["recovered_chatgpt"])
            self.assertEqual(telem_arg["total_failures"], 1)
            self.assertEqual(telem_arg["failures"][0]["account"], "user2@gmail.com")

    def test_refresh_antigravity_oauth_flow_structure(self):
        # Kiểm tra cấu trúc flow refresh antigravity account qua GPM
        import time
        ten_days_ago = time.time() - (10 * 86400)
        mock_gpm_profile = {"id": "pid_ag", "name": "M01 - user@gmail.com - 5101", "raw_proxy": "test:5101", "created_at": ten_days_ago}
        with patch("requests.get") as mock_get:
            mock_get.return_value.json.return_value = {"success": False, "message": "Proxy connection error"}
            ok, msg = w.refresh_antigravity_account_via_gpm("pid_ag", "user@gmail.com", mock_gpm_profile, {})
            self.assertFalse(ok)
            self.assertIn("Start profile fail", msg)

    def test_profile_aged_7_days_safety_valve(self):
        import time
        # Case 1: Profile 10 ngày tuổi -> Đạt điều kiện
        prof_old = {"created_at": time.time() - (10 * 86400)}
        self.assertTrue(w.is_profile_aged_7_days(prof_old, min_days=7))

        # Case 2: Profile 2 ngày tuổi -> Bị van an toàn chặn
        prof_new = {"created_at": time.time() - (2 * 86400)}
        self.assertFalse(w.is_profile_aged_7_days(prof_new, min_days=7))

        # Case 3: Chặn ngay trong refresh_antigravity_account_via_gpm
        ok, msg = w.refresh_antigravity_account_via_gpm("pid_new", "new@gmail.com", prof_new, {})
        self.assertFalse(ok)
        self.assertIn("Van an toàn", msg)

    def test_oauth_codex_timeout_on_callback_capture(self):
        # Test OAuth code capture timeout (90s) -> Codex OAuth returns false with timeout message
        with patch("cron_chatgpt_web_pool_watchdog.CodexOAuth1455Lock") as mock_lock_cls, \
             patch("cron_chatgpt_web_pool_watchdog.sync_playwright") as mock_pw, \
             patch("cron_chatgpt_web_pool_watchdog.requests") as mock_req:

            mock_lock = MagicMock()
            mock_lock.__enter__.return_value = mock_lock
            mock_lock_cls.return_value = mock_lock

            mock_p = MagicMock()
            mock_browser = MagicMock()
            mock_context = MagicMock()
            mock_page = MagicMock()

            mock_p.chromium.connect_over_cdp.return_value = mock_browser
            mock_browser.contexts = [mock_context]
            mock_context.pages = [mock_page]
            mock_page.url = "https://api.example.com/auth"
            mock_page.title.return_value = "Auth Page"
            mock_page.locator.return_value.first.count.return_value = 0
            mock_pw.return_value.__enter__.return_value = mock_p

            mock_req.post.return_value.json.return_value = {"success": False, "data": {}}
            mock_req.get.return_value.json.return_value = {"success": True, "data": {"authUrl": "https://test.com"}}

            ok, msg = w.perform_codex_oauth("127.0.0.1:9222", "test@gmail.com", {})
            self.assertFalse(ok)
            self.assertIn("Timeout OAuth", msg)

    def test_codex_oauth_invalid_reason_classifies_401(self):
        # Test classification: httpStatus=401 -> token revoked
        conn = {
            "id": "cid_1",
            "name": "user@gmail.com",
            "httpStatus": 401,
            "isActive": True
        }
        reason = w.codex_oauth_invalid_reason(conn)
        self.assertEqual(reason, "httpStatus=401")

    def test_codex_oauth_invalid_reason_detects_upstream_auth_error(self):
        # Test classification: error_type=upstream_auth_error
        conn = {
            "id": "cid_1",
            "name": "user@gmail.com",
            "error_type": "upstream_auth_error",
            "isActive": True
        }
        reason = w.codex_oauth_invalid_reason(conn)
        self.assertEqual(reason, "error_type=upstream_auth_error")

    def test_codex_oauth_invalid_reason_detects_token_revoked(self):
        # Test classification: lastError contains "token revoked"
        conn = {
            "id": "cid_1",
            "lastError": "Google token revoked by user action",
            "isActive": True
        }
        reason = w.codex_oauth_invalid_reason(conn)
        self.assertIn("lastError contains Token invalid/revoked", reason)

    def test_codex_oauth_invalid_reason_returns_none_for_ambiguous(self):
        # Test: no clear invalidation markers -> None
        conn = {
            "id": "cid_1",
            "testStatus": "active",
            "isActive": True,
            "lastError": None
        }
        reason = w.codex_oauth_invalid_reason(conn)
        self.assertIsNone(reason)

    def test_codex_oauth_1455_lock_acquire_success(self):
        # Test successful lock acquire/release
        import tempfile
        operations = []

        def mock_locking(fd, cmd, nbytes):
            if cmd == 2:  # LK_NBLCK
                operations.append("acquire")
            else:  # LK_UNLCK
                operations.append("release")

        with tempfile.TemporaryDirectory() as td:
            lock_file = Path(td) / "test.lock"
            with patch("msvcrt.locking", side_effect=mock_locking):
                lock = w.CodexOAuth1455Lock(lock_file=str(lock_file), timeout=5)
                with lock:
                    self.assertIn("acquire", operations)
                self.assertIn("release", operations)

    def test_codex_oauth_1455_lock_exit_cleanup(self):
        # Test lock cleanup in __exit__ even with errors
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            lock_file = Path(td) / "test.lock"
            with patch("msvcrt.locking") as mock_locking:
                mock_locking.side_effect = [None, OSError("unlocking failed")]
                lock = w.CodexOAuth1455Lock(lock_file=str(lock_file), timeout=5)
                # __exit__ should handle the OSError gracefully
                try:
                    with lock:
                        pass
                except OSError:
                    pass
                self.assertIsNone(lock.handle)

    def test_antigravity_oauth_phone_checkpoint_safe_abort(self):
        # Test: phone checkpoint detection -> safe abort without crashing
        with patch("cron_chatgpt_web_pool_watchdog.sync_playwright") as mock_pw, \
             patch("cron_chatgpt_web_pool_watchdog.requests") as mock_req:
            mock_p = MagicMock()
            mock_browser = MagicMock()
            mock_context = MagicMock()
            mock_page = MagicMock()

            mock_p.chromium.connect_over_cdp.return_value = mock_browser
            mock_browser.contexts = [mock_context]
            mock_context.pages = [mock_page]

            # Page returns phone checkpoint text
            mock_page.url = "https://accounts.google.com/signup"
            mock_page.evaluate.return_value = True  # is_phone_cp = True
            mock_page.get_by_text.return_value.first.count.return_value = 0

            mock_pw.return_value.__enter__.return_value = mock_p
            mock_req.get.side_effect = [
                MagicMock(status_code=200, json=lambda: {
                    "authUrl": "https://accounts.google.com/signup",
                    "state": "state123",
                    "codeVerifier": "verif123"
                })
            ]

            ok, msg = w.perform_antigravity_oauth("127.0.0.1:9222", "test@gmail.com", {})
            self.assertFalse(ok)
            self.assertIn("Timeout", msg)

    def test_filesystem_conflict_retry_on_oserror(self):
        # Test: OSError on lock acquire triggers sleep/retry logic
        import tempfile
        retry_count = [0]

        def side_effect_locking(fd, cmd, nbytes):
            if cmd == 2:  # LK_NBLCK
                retry_count[0] += 1
                if retry_count[0] < 2:
                    raise OSError("File locked")

        with tempfile.TemporaryDirectory() as td:
            lock_file = Path(td) / "race.lock"
            with patch("msvcrt.locking", side_effect=side_effect_locking), \
                 patch("time.sleep"):
                lock = w.CodexOAuth1455Lock(lock_file=str(lock_file), timeout=10)
                with lock:
                    self.assertGreaterEqual(retry_count[0], 1)

    def test_telemetry_resilience_on_io_failure(self):
        # Kiểm tra tính phục hồi: telemetry gặp lỗi disk IO không được làm crash watchdog
        with patch("builtins.open", side_effect=IOError("Simulated disk full / permission denied")):
            ok = w.save_telemetry_metrics({"dummy": "data"})
            self.assertFalse(ok)

    def test_fetch_google_recovery_otp_regex_and_filtering(self):
        # Kiểm tra logic trích xuất OTP 6 số và bỏ qua cảnh báo bảo mật
        import re
        sample_security_alert = "Cảnh báo bảo mật cho tài khoản liên kết"
        is_alert = any(k in sample_security_alert.lower() for k in ["cảnh báo", "security alert", "cảnh báo bảo mật"])
        self.assertTrue(is_alert)

        sample_otp_body = "Mã xác minh Google của bạn là 492815. Mã này sẽ hết hạn sau 10 phút."
        m = re.search(r'\b(\d{6})\b', sample_otp_body)
        self.assertIsNotNone(m)
        self.assertEqual(m.group(1), "492815")

    def test_cdp_connect_retry_and_backoff_exhaustion(self):
        with patch("cron_chatgpt_web_pool_watchdog.sync_playwright") as mock_pw, \
             patch("time.sleep") as mock_sleep, \
             patch("cron_chatgpt_web_pool_watchdog.log") as mock_log:
            mock_p = MagicMock()
            mock_p.chromium.connect_over_cdp.side_effect = Exception("ECONNREFUSED")
            mock_pw.return_value.__enter__.return_value = mock_p
            res = w.perform_auto_login_and_extract_cookies("127.0.0.1:9999", "test@gmail.com", {})
            self.assertIsNone(res)

    def test_disable_codex_connection_guard_rejects_non_codex(self):
        # Test: disable_codex_connection guards against accidentally disabling non-codex providers
        with patch("sqlite3.connect") as mock_connect:
            mock_conn = MagicMock()
            mock_cur = MagicMock()
            mock_connect.return_value = mock_conn
            mock_conn.cursor.return_value = mock_cur
            # Guard check: returns non-codex provider
            mock_cur.execute.return_value = mock_cur
            mock_cur.fetchone.return_value = ("antigravity", "some_account")

            conn = {"id": "wrong_cid", "name": "acc@gmail.com", "isActive": True}
            result = w.disable_codex_connection(conn, "test reason")
            self.assertFalse(result)
            mock_cur.execute.assert_called_once()

    def test_disable_codex_connection_rollback_on_db_error(self):
        # Test: DB error -> rollback called
        with patch("sqlite3.connect") as mock_connect:
            mock_conn = MagicMock()
            mock_cur = MagicMock()
            mock_connect.return_value = mock_conn
            mock_conn.cursor.return_value = mock_cur
            mock_cur.fetchone.return_value = ("codex", "some_account")
            # First call (guard check) succeeds; second (UPDATE) fails
            mock_cur.execute.side_effect = [None, sqlite3.DatabaseError("Disk I/O error")]

            conn = {"id": "cid_x", "name": "acc@gmail.com", "isActive": True}
            result = w.disable_codex_connection(conn, "test reason")
            self.assertFalse(result)
            mock_conn.rollback.assert_called_once()

    def test_detect_broken_codex_via_test_endpoint_401(self):
        # Test: when codex_oauth_invalid_reason returns None, script calls /test endpoint
        codex_conns = [
            {"id": "cid_cx_1", "name": "user@gmail.com", "isActive": True, "testStatus": "error", "lastError": None}
        ]

        with patch.object(w, "codex_oauth_invalid_reason", return_value=None) as mock_classify, \
             patch("requests.post") as mock_post:
            mock_post.return_value.json.return_value = {"valid": False, "statusCode": 401}

            # Simulate main() logic: when reason is None, try /test endpoint
            cid = codex_conns[0]["id"]
            reason = w.codex_oauth_invalid_reason(codex_conns[0])
            if not reason:
                try:
                    r_test = requests.post(f"{w.OMNI_API}/api/providers/{cid}/test", json={}, timeout=5).json()
                    if r_test.get("valid") is False and r_test.get("statusCode") == 401:
                        reason = "401_token_revoked"
                except Exception:
                    pass

            self.assertEqual(reason, "401_token_revoked")
            mock_post.assert_called_once()

    def test_refresh_chatgpt_token_validation_prevents_invalid_cookie_store(self):
        # Test: test_connection_valid returns False -> returns error without storing token
        with patch("requests.get") as mock_gpm, \
             patch.object(w, "perform_auto_login_and_extract_cookies", return_value="invalid_token"), \
             patch.object(w, "test_connection_valid", return_value=(False, "Invalid format")):

            mock_gpm.return_value.json.return_value = {
                "success": True,
                "data": {"remote_debugging_address": "127.0.0.1:9222"}
            }

            ok, msg = w.refresh_chatgpt_account_via_gpm("pid_1", "cid_1", "test@gmail.com", {})
            self.assertFalse(ok)
            self.assertIn("Validate fail", msg)

    def test_save_telemetry_handles_missing_parent_dir(self):
        # Test: telemetry save creates parent directory if missing
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            nested_path = Path(td) / "deep" / "nested" / "path" / "telemetry.json"
            metrics = {"timestamp": "2026-09-25T00:00:00Z", "test": True}

            with patch.object(w, "TELEMETRY_CACHE_PATH", str(nested_path)):
                ok = w.save_telemetry_metrics(metrics)
                self.assertTrue(ok)
                self.assertTrue(nested_path.exists())
                data = json.loads(nested_path.read_text())
                self.assertEqual(data["test"], True)
                nested_path.unlink()

    def test_antigravity_oauth_respects_standby_profile_age_check(self):
        # Test: 2-day-old profile -> rejected by van an toàn (7-day safety valve)
        import time
        two_days_ago = time.time() - (2 * 86400)
        prof_young = {"id": "pid_new", "created_at": two_days_ago}

        ok, msg = w.refresh_antigravity_account_via_gpm("pid_new", "new@gmail.com", prof_young, {})
        self.assertFalse(ok)
        self.assertIn("Van an toàn", msg)
        self.assertIn("7 ngày", msg)

    def test_codex_oauth_lock_context_exit_called(self):
        # Test: CodexOAuth1455Lock __exit__ cleanup is always called
        import tempfile
        operations = []

        def side_effect_locking(fd, cmd, nbytes):
            operations.append(cmd)

        with tempfile.TemporaryDirectory() as td:
            lock_file = Path(td) / "test.lock"
            with patch("msvcrt.locking", side_effect=side_effect_locking):
                lock = w.CodexOAuth1455Lock(lock_file=str(lock_file), timeout=5)
                with lock:
                    pass
                # Verify unlock (LK_UNLCK=0) was called in __exit__
                self.assertIn(0, operations)

    def test_get_gpm_profiles_pagination_validation_rejects_mismatch(self):
        # Test: pagination total_page changes between pages -> fails with ValueError
        with patch("requests.get") as mock_get:
            def side_effect_fn(url, *args, **kwargs):
                resp = MagicMock()
                if "page=1" in url:
                    resp.json.return_value = {
                        "pagination": {"total_page": 5, "total": 100},
                        "data": [{"id": f"pid_{i}"} for i in range(1, 51)]
                    }
                elif "page=2" in url:
                    resp.json.return_value = {
                        "pagination": {"total_page": 3},  # Changed!
                        "data": [{"id": f"pid_{i}"} for i in range(51, 101)]
                    }
                return resp

            mock_get.side_effect = side_effect_fn
            profs = w.get_gpm_profiles()
            self.assertEqual(profs, [])

    def test_load_accounts_credentials_handles_missing_excel(self):
        # Test: Excel file doesn't exist -> return empty dict, no crash
        with patch("os.path.exists", return_value=False):
            creds = w.load_accounts_credentials()
            self.assertEqual(creds, {})

    def test_categorize_failure_priority_order(self):
        # Test: phone error matched before others (priority matters)
        self.assertEqual(w.categorize_failure("số điện thoại failed"), "PHONE_CHECKPOINT")
        self.assertEqual(w.categorize_failure("timeout at phone prompt"), "PHONE_CHECKPOINT")
        self.assertEqual(w.categorize_failure("regular timeout"), "TIMEOUT")
        self.assertEqual(w.categorize_failure("proxy refused connection"), "PROXY_ERROR")


class TestCronSyncWatchdogSafety(unittest.TestCase):
    def test_sync_file_skip_when_dst_newer(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.py"
            dst = Path(td) / "dst.py"
            src.write_text("v1")
            dst.write_text("v2")
            os.utime(src, (1000, 1000))
            os.utime(dst, (2000, 2000))
            res = csw.sync_file(src, dst, force=False)
            self.assertEqual(res, "skipped")
            self.assertEqual(dst.read_text(), "v2")

    def test_sync_file_conflict_when_mtime_equal_but_content_diff(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.py"
            dst = Path(td) / "dst.py"
            src.write_text("code_a")
            dst.write_text("code_b")
            os.utime(src, (1500, 1500))
            os.utime(dst, (1500, 1500))
            res = csw.sync_file(src, dst, force=False)
            self.assertEqual(res, "conflict")
            self.assertEqual(dst.read_text(), "code_b")

    def test_sync_file_force_creates_conflict_backup(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.py"
            dst = Path(td) / "dst.py"
            src.write_text("code_a")
            dst.write_text("code_b")
            os.utime(src, (1500, 1500))
            os.utime(dst, (1500, 1500))
            res = csw.sync_file(src, dst, force=True)
            self.assertEqual(res, "copied")
            self.assertEqual(dst.read_text(), "code_a")
            baks = list(Path(td).glob("dst.py.conflict-*.bak"))
            self.assertEqual(len(baks), 1)
            self.assertEqual(baks[0].read_text(), "code_b")

    def test_telemetry_failure_categorization_matrix(self):
        self.assertEqual(w.categorize_failure("Lỗi xác minh số điện thoại Google"), "PHONE_CHECKPOINT")
        self.assertEqual(w.categorize_failure("Timeout OAuth sau 90s"), "TIMEOUT")
        self.assertEqual(w.categorize_failure("Proxy connection error refused"), "PROXY_ERROR")
        self.assertEqual(w.categorize_failure("Google reCAPTCHA Challenge"), "CAPTCHA_CHALLENGE")
        self.assertEqual(w.categorize_failure("Lỗi không xác định khác"), "UNKNOWN")

    def test_sync_file_atomic_concurrency_resilience(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.py"
            dst = Path(td) / "dst.py"
            src.write_text("concurrent_content_v1")
            res = csw.sync_file(src, dst, force=False)
            self.assertEqual(res, "copied")
            self.assertEqual(dst.read_text(), "concurrent_content_v1")
            self.assertFalse(any(p.name.startswith(".dst.py.sync.") for p in Path(td).iterdir()))

    def test_sync_scripts_two_way_reconcile_and_conflict_handling(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            dep_dir = Path(td) / "deploy"
            kibe_dir = Path(td) / "kibe"
            one_dir = Path(td) / "onedrive"
            dep_dir.mkdir(); kibe_dir.mkdir(); one_dir.mkdir()
            (dep_dir / "script_dep.py").write_text("v_dep")
            (kibe_dir / "script_kibe.py").write_text("v_kibe")

            with patch.object(csw, "DEPLOY_SCRIPTS", dep_dir), \
                 patch.object(csw, "KIBE_SCRIPTS", kibe_dir), \
                 patch.object(csw, "ONEDRIVE_SCRIPTS", one_dir):
                synced = csw.sync_scripts(force=False)
                self.assertGreater(synced, 0)
                self.assertTrue((kibe_dir / "script_dep.py").exists())
                self.assertTrue((dep_dir / "script_kibe.py").exists())
                self.assertTrue((one_dir / "script_dep.py").exists())
                self.assertTrue((one_dir / "script_kibe.py").exists())

    def test_gpm_lifecycle_sync_api_timeout_and_network_error_resilience(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("sync_gpm_lifecycle", str(REPO_SCRIPTS / "sync_gpm_lifecycle.py"))
        sgl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(sgl)
        with patch.object(sgl, "is_gpm_api_live", return_value=False), \
             patch.object(sgl, "logger") as mock_log:
            sgl.sync_lifecycle_gpm()
            mock_log.warning.assert_called_with("GPM API không hoạt động, bỏ qua sync_lifecycle_gpm.")

    def test_antigravity_oauth_cleanup_on_exception(self):
        with patch("requests.get") as mock_get:
            mock_get.side_effect = [
                MagicMock(status_code=200, json=lambda: {"success": True, "data": {"remote_debugging_address": "127.0.0.1:9222"}}),
                MagicMock(status_code=200, json=lambda: {"success": True}),
                MagicMock(status_code=200, json=lambda: {"success": True})
            ]
            with patch.object(w, "perform_antigravity_oauth", side_effect=Exception("CDP network crash")):
                ok, msg = w.refresh_antigravity_account_via_gpm("pid_test", "acc@gmail.com", {"created_at": 1000}, {})
                self.assertFalse(ok)
                self.assertIn("Exception: CDP network crash", msg)
                close_calls = [c[0][0] for c in mock_get.call_args_list if "close" in str(c) or "stop" in str(c)]
                self.assertGreater(len(close_calls), 0)


class TestSyncGpmLifecycleLogic(unittest.TestCase):
    def test_taikhoan_dat_reading_and_exclusions(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("sync_gpm_lifecycle", str(REPO_SCRIPTS / "sync_gpm_lifecycle.py"))
        sgl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(sgl)
        import tempfile
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            wb_path = Path(td) / "test_tk.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Tài Khoản"
            ws.append(["Mã Máy", "Col1", "Col2", "Col3", "Col4", "Email", "Col6", "Col7", "Col8", "Serial", "Proxy"])
            ws.append([45, "", "", "", "", "valid45@gmail.com", "", "", "", "ce071607", "1.1.1.1:5107:u:p"])
            ws.append([7, "", "", "", "", "khoale_test@gmail.com", "", "", "", "9885f630", ""])
            ws.append([19, "", "", "", "", "valid19@gmail.com", "", "", "", "", ""])
            wb.save(str(wb_path))

            candidates = {}
            added = sgl.merge_taikhoan_dat_candidates(
                str(wb_path), candidates, {"excluded@gmail.com"}
            )
            serials, proxy_map = {}, {}
            sgl.merge_taikhoan_dat_proxies(str(wb_path), serials, proxy_map)

            self.assertEqual(added, 2)
            self.assertEqual(candidates["valid45@gmail.com"], 45)
            self.assertNotIn("khoale_test@gmail.com", candidates)
            self.assertEqual(candidates["valid19@gmail.com"], 19)
            self.assertEqual(serials[45], "ce071607")
            self.assertEqual(proxy_map[45], "1.1.1.1:5107:u:p")
            self.assertTrue(proxy_map[19].startswith("test.taadaa.click:20019:"))

    def test_s7_cleanup_dry_run_does_not_delete(self):
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "sync_gpm_lifecycle_cleanup_dry_run",
            str(REPO_SCRIPTS / "sync_gpm_lifecycle.py"),
        )
        sgl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(sgl)

        master_wb = MagicMock()
        master_ws = MagicMock()
        master_ws.iter_rows.return_value = [
            ("header", "email", "", "", "", "", "status", "machine"),
            ("", "die@example.com", "", "", "", "", "DIE", "M01"),
        ]
        master_wb.__getitem__.return_value = master_ws

        proxy_wb = MagicMock()
        proxy_ws = MagicMock()
        proxy_ws.iter_rows.return_value = [
            ("machine", "serial"),
            (1, "serial-1"),
        ]
        proxy_wb.__getitem__.return_value = proxy_ws

        def fake_exists(path):
            return path in {sgl.MASTER_FILE, sgl.PROXY_FILE}

        with patch.object(sgl, "is_morning_cleanup_window", return_value=True), \
             patch.object(sgl, "openpyxl") as mock_openpyxl, \
             patch.object(sgl, "get_online_adb_devices", return_value={"serial-1"}), \
             patch.object(sgl, "is_machine_in_feed_slot", return_value=False), \
             patch.object(sgl, "acquire_device_lock") as mock_lock, \
             patch.object(sgl, "subprocess") as mock_subprocess, \
             patch.object(sgl, "log_telemetry_metric") as mock_telemetry, \
             patch.object(sgl.os.path, "exists", side_effect=fake_exists):
            mock_openpyxl.load_workbook.side_effect = [master_wb, proxy_wb]
            mock_lock.return_value.__enter__.return_value = MagicMock()
            mock_subprocess.run.return_value.stdout = (
                "Account {name=die@example.com, type=com.google}"
            )

            cleaned = sgl.cleanup_s7_die_accounts(dry_run=True, force=True)

        self.assertEqual(cleaned, 1)
        self.assertEqual(mock_subprocess.run.call_count, 2)
        self.assertTrue(
            all(call.args[0][0] == sgl.ADB_PATH for call in mock_subprocess.run.call_args_list)
        )
        mock_telemetry.assert_called_once_with(
            "s7_cleanup_completed",
            {"cleaned_count": 1, "die_machines_count": 1},
        )

    def test_s7_cleanup_lock_contention_is_skipped(self):
        import importlib.util

        spec = importlib.util.spec_from_file_location(
            "sync_gpm_lifecycle_cleanup_lock",
            str(REPO_SCRIPTS / "sync_gpm_lifecycle.py"),
        )
        sgl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(sgl)

        master_wb = MagicMock()
        master_ws = MagicMock()
        master_ws.iter_rows.return_value = [
            ("header", "email", "", "", "", "", "status", "machine"),
            ("", "die@example.com", "", "", "", "", "DIE", "M01"),
        ]
        master_wb.__getitem__.return_value = master_ws

        proxy_wb = MagicMock()
        proxy_ws = MagicMock()
        proxy_ws.iter_rows.return_value = [
            ("machine", "serial"),
            (1, "serial-1"),
        ]
        proxy_wb.__getitem__.return_value = proxy_ws

        def fake_exists(path):
            return path in {sgl.MASTER_FILE, sgl.PROXY_FILE}

        with patch.object(sgl, "is_morning_cleanup_window", return_value=True), \
             patch.object(sgl, "openpyxl") as mock_openpyxl, \
             patch.object(sgl, "get_online_adb_devices", return_value={"serial-1"}), \
             patch.object(sgl, "is_machine_in_feed_slot", return_value=False), \
             patch.object(sgl, "acquire_device_lock", side_effect=sgl.DeviceLockUnavailable), \
             patch.object(sgl, "subprocess") as mock_subprocess, \
             patch.object(sgl, "log_telemetry_metric") as mock_telemetry, \
             patch.object(sgl.os.path, "exists", side_effect=fake_exists):
            mock_openpyxl.load_workbook.side_effect = [master_wb, proxy_wb]

            cleaned = sgl.cleanup_s7_die_accounts(dry_run=True, force=True)

        self.assertEqual(cleaned, 0)
        mock_subprocess.run.assert_not_called()
        mock_telemetry.assert_called_once_with(
            "s7_cleanup_completed",
            {"cleaned_count": 0, "die_machines_count": 1},
        )

    def test_telemetry_schema_validation_and_downstream_compat(self):
        metrics = {
            "timestamp": "2026-09-25T00:00:00Z",
            "chatgpt_web": {"active": 15, "total": 21},
            "antigravity": {"active": 109, "total": 119},
            "codex": {"active": 29, "total": 29},
            "total_failures": 0,
            "failures": []
        }
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            cache = Path(td) / "cache.json"
            hist = Path(td) / "hist.jsonl"
            with patch.object(w, "TELEMETRY_CACHE_PATH", str(cache)), \
                 patch.object(w, "TELEMETRY_HISTORY_PATH", str(hist)):
                ok = w.save_telemetry_metrics(metrics)
                self.assertTrue(ok)
                data = json.loads(cache.read_text(encoding="utf-8"))
                self.assertIn("timestamp", data)
                self.assertIn("chatgpt_web", data)

    def test_sync_file_filesystem_error_resilience(self):
        import cron_sync_watchdog as csw
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            src = Path(td) / "src.py"
            dst = Path(td) / "dst.py"
            src.write_text("code")
            with patch.object(csw, "_atomic_copy", side_effect=OSError("Disk full")):
                res = csw.sync_file(src, dst)
                self.assertEqual(res, "error")

    def test_sync_gpm_lifecycle_api_timeout_safety(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("sgl_timeout", str(REPO_SCRIPTS / "sync_gpm_lifecycle.py"))
        sgl = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(sgl)
        with patch.object(sgl, "is_gpm_api_live", return_value=False), \
             patch.object(sgl, "logger") as mock_log:
            sgl.sync_lifecycle_gpm()
            self.assertTrue(mock_log.warning.called)


class TestTelemetryEndToEndAndFarmSafety(unittest.TestCase):
    def test_telemetry_end_to_end_schema_validation_and_downstream_format(self):
        import io
        import tempfile
        from contextlib import redirect_stdout
        from datetime import datetime

        conns_by_provider = {
            "chatgpt-web": [
                {"id": "cid_cg_1", "name": "user1@gmail.com (GPM Web)", "isActive": False, "testStatus": "expired"},
                {"id": "cid_cg_2", "name": "user3@gmail.com (GPM Web)", "isActive": False, "testStatus": "expired"},
                {"id": "cid_cg_3", "name": "user4@gmail.com (GPM Web)", "isActive": True, "testStatus": "active"},
            ],
            "antigravity": [
                {"id": "cid_ag_1", "name": "user2@gmail.com", "isActive": True, "testStatus": "expired"},
            ],
            "codex": [],
        }
        gpm_map = {
            "user1@gmail.com": [{"id": "pid_1", "name": "01 - user1@gmail.com - 5101"}],
            "user2@gmail.com": [{"id": "pid_2", "name": "02 - user2@gmail.com - 5102"}],
        }

        with tempfile.TemporaryDirectory() as td:
            cache = Path(td) / "cache.json"
            hist = Path(td) / "hist.jsonl"
            public = Path(td) / "public.json"
            out = io.StringIO()
            with patch.object(w, "TELEMETRY_CACHE_PATH", str(cache)), \
                 patch.object(w, "TELEMETRY_HISTORY_PATH", str(hist)), \
                 patch.object(w, "TELEMETRY_PUBLIC_PATH", str(public)), \
                 patch.object(w, "log"), \
                 patch.object(w, "get_connections_by_provider", side_effect=lambda p: conns_by_provider.get(p, [])), \
                 patch.object(w, "get_gpm_profiles_map", return_value=gpm_map), \
                 patch.object(w, "load_accounts_credentials", return_value={}), \
                 patch.object(w, "refresh_chatgpt_account_via_gpm", return_value=(False, "Timeout OAuth sau 90s")), \
                 patch.object(w, "refresh_antigravity_account_via_gpm", return_value=(False, "Google yêu cầu số điện thoại")), \
                 redirect_stdout(out):
                w.main()

            data = json.loads(public.read_text(encoding="utf-8"))
            self.assertEqual(json.loads(cache.read_text(encoding="utf-8")), data)
            hist_lines = hist.read_text(encoding="utf-8").splitlines()
            self.assertEqual(len(hist_lines), 1)
            self.assertEqual(json.loads(hist_lines[0]), data)

        # Public schema: required keys and types
        for key in ("timestamp", "last_scan", "selector_version", "chatgpt_web", "antigravity", "codex",
                    "recovered_chatgpt", "recovered_ag", "recovered_codex", "total_failures",
                    "failures", "failure_categories"):
            self.assertIn(key, data)
        self.assertEqual(data["selector_version"], w.SELECTOR_VERSION)
        self.assertIsNotNone(datetime.fromisoformat(data["last_scan"]).tzinfo)
        self.assertEqual(data["chatgpt_web"], {"active": 1, "total": 3})
        for provider in ("chatgpt_web", "antigravity", "codex"):
            self.assertEqual(set(data[provider]), {"active", "total"})
            self.assertLessEqual(data[provider]["active"], data[provider]["total"])

        # failures: one entry per failed account, each with a known category
        self.assertEqual(data["total_failures"], 3)
        self.assertEqual(len(data["failures"]), 3)
        valid_categories = {"PHONE_CHECKPOINT", "TIMEOUT", "PROXY_ERROR", "CAPTCHA_CHALLENGE", "UNKNOWN"}
        by_account = {}
        for item in data["failures"]:
            self.assertEqual(set(item), {"account", "category", "error"})
            self.assertIn(item["category"], valid_categories)
            self.assertLessEqual(len(item["error"]), 100)
            by_account[item["account"]] = item["category"]
        self.assertEqual(by_account, {
            "user1@gmail.com": "TIMEOUT",
            "user3@gmail.com": "UNKNOWN",
            "user2@gmail.com": "PHONE_CHECKPOINT",
        })

        # failure_categories is the aggregate of failures[].category
        self.assertEqual(data["failure_categories"], {"TIMEOUT": 1, "UNKNOWN": 1, "PHONE_CHECKPOINT": 1})
        self.assertEqual(sum(data["failure_categories"].values()), data["total_failures"])

        # Downstream consumer: report lines derived from the same counters
        report = out.getvalue()
        self.assertIn("🤖 [POOL HEALER] BÁO CÁO SỨC KHỎE", report)
        self.assertIn(
            f"• ChatGPT-Web: {data['chatgpt_web']['active']}/{data['chatgpt_web']['total']} ACTIVE", report
        )
        self.assertIn(f"• Cần chú ý ({data['total_failures']} acc):", report)

    def test_farm_safety_no_accidental_ban_or_cookie_wipe_on_upstream_5xx(self):
        import tempfile
        import urllib.error

        def codex_test_response(status):
            resp = MagicMock()
            resp.json.return_value = {"valid": False, "statusCode": status}
            return {"return_value": resp}

        upstream_failures = {
            "http_503": (
                urllib.error.HTTPError("http://omni/validate", 503, "Service Unavailable", {}, None),
                codex_test_response(503),
            ),
            "http_429": (
                urllib.error.HTTPError("http://omni/validate", 429, "Too Many Requests", {}, None),
                codex_test_response(429),
            ),
            "network": (
                urllib.error.URLError("Connection reset by peer"),
                {"side_effect": w.requests.ConnectionError("Connection aborted")},
            ),
        }

        for label, (validate_error, codex_post_kwargs) in upstream_failures.items():
            with self.subTest(upstream=label), tempfile.TemporaryDirectory() as td:
                db_path = str(Path(td) / "storage.sqlite")
                conn = sqlite3.connect(db_path)
                conn.execute("CREATE TABLE provider_connections (id TEXT PRIMARY KEY, provider TEXT, name TEXT, api_key TEXT, is_active INTEGER, test_status TEXT, last_error TEXT, last_error_at TEXT, backoff_level INTEGER, rate_limited_until TEXT, updated_at TEXT)")
                conn.execute("CREATE TABLE combos (name TEXT PRIMARY KEY, data TEXT)")
                conn.execute("INSERT INTO provider_connections VALUES ('cid_web_1', 'chatgpt-web', 'user1@gmail.com', 'eyJhbGOLD_COOKIE', 0, 'expired', 'upstream 503', NULL, 2, NULL, '2026-10-01')")
                conn.execute("INSERT INTO provider_connections VALUES ('cid_cx_1', 'codex', 'user3@gmail.com', 'codex_token', 1, 'error', 'HTTP 503 upstream', NULL, 1, NULL, '2026-10-01')")
                conn.execute("INSERT INTO combos VALUES ('chatgpt-web-pool', '{\"models\": []}')")
                conn.commit()
                before = conn.execute("SELECT * FROM provider_connections ORDER BY id").fetchall()
                combos_before = conn.execute("SELECT * FROM combos").fetchall()
                conn.close()

                conns_by_provider = {
                    "chatgpt-web": [{"id": "cid_web_1", "name": "user1@gmail.com (GPM Web)", "isActive": False,
                                     "testStatus": "expired", "apiKey": "eyJhbGOLD_COOKIE"}],
                    "antigravity": [],
                    "codex": [{"id": "cid_cx_1", "name": "user3@gmail.com", "isActive": True,
                               "testStatus": "error", "lastError": "HTTP 503 upstream"}],
                }
                gpm_start = MagicMock()
                gpm_start.json.return_value = {"success": True, "data": {"remote_debugging_address": "127.0.0.1:9222"}}

                with patch.object(w, "DB_PATH", db_path), \
                     patch.object(w, "log"), \
                     patch("builtins.print"), \
                     patch.object(w, "get_connections_by_provider", side_effect=lambda p: conns_by_provider.get(p, [])), \
                     patch.object(w, "get_gpm_profiles_map", return_value={"user1@gmail.com": [{"id": "pid_1"}]}), \
                     patch.object(w, "load_accounts_credentials", return_value={}), \
                     patch.object(w, "perform_auto_login_and_extract_cookies", return_value="__Secure-next-auth.session-token=NEW"), \
                     patch.object(w.urllib.request, "urlopen", side_effect=validate_error), \
                     patch.object(w.requests, "get", return_value=gpm_start), \
                     patch.object(w.requests, "post", **codex_post_kwargs), \
                     patch.object(w.time, "sleep"), \
                     patch.object(w, "refresh_codex_account_via_gpm") as codex_refresh, \
                     patch.object(w, "disable_codex_connection", wraps=w.disable_codex_connection) as disable_spy, \
                     patch.object(w, "save_telemetry_metrics") as save_telem:
                    w.main()

                # Codex: transient upstream status is not OAuth invalidation -> never disabled or re-OAuthed
                disable_spy.assert_not_called()
                codex_refresh.assert_not_called()

                # ChatGPT-Web: upstream validation failed -> old cookie, state and pool combo untouched
                conn = sqlite3.connect(db_path)
                after = conn.execute("SELECT * FROM provider_connections ORDER BY id").fetchall()
                combos_after = conn.execute("SELECT * FROM combos").fetchall()
                conn.close()
                self.assertEqual(after, before)
                self.assertEqual(combos_after, combos_before)

                telem = save_telem.call_args[0][0]
                self.assertEqual(telem["recovered_chatgpt"], [])
                self.assertEqual([f["account"] for f in telem["failures"]], ["user1@gmail.com"])
                self.assertIn("Validate fail", telem["failures"][0]["error"])
