#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scan_two_farm_avatar_duplicates.py
Công cụ quét đối soát toàn diện mã băm MD5 của avatar.jpg trên toàn bộ 2 Farm (Kibe & Admin).
Tự động đối chiếu với các workbook để xác định chính xác số máy và tài khoản bị trùng lặp.

Usage:
    python scan_two_farm_avatar_duplicates.py [--json-out PATH]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

ROOT_KIBE_NUOI = Path(r"D:\TIKTOK-videonuoinick")
ROOT_KIBE_GOC = Path(r"D:\video goc")
ROOT_ADMIN_GOC = Path(r"D:\video goc may 2")

EXCEL_KIBE_SAFE = Path(r"D:\OneDrive\TaadaaData\kibe\taikhoan_run_safe.xlsx")
EXCEL_ADMIN_MASTER = Path(r"D:\OneDrive\TaadaaData\admin\taikhoan_dat_v2_updated .xlsx")
DIR_KIBE_WORKBOOKS = Path(r"D:\OneDrive\TaadaaData\kibe")
DIR_ADMIN_WORKBOOKS = Path(r"D:\OneDrive\TaadaaData\admin")


def get_md5(p: Path) -> str:
    h = hashlib.md5()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def load_account_mappings() -> dict[str, list[dict[str, Any]]]:
    """Load mapping từ folder_num -> danh sách thông tin account."""
    folder_map: dict[str, list[dict[str, Any]]] = defaultdict(list)
    try:
        import openpyxl
    except ImportError:
        print("[WARN] Thiếu thư viện openpyxl, bỏ qua bước map tên account.")
        return folder_map

    # 1. Kibe taikhoan_run_safe
    if EXCEL_KIBE_SAFE.exists():
        try:
            wb = openpyxl.load_workbook(EXCEL_KIBE_SAFE, data_only=True, read_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 3:
                    continue
                m, f, acc = row[0], row[1], row[2]
                name = row[3] if len(row) > 3 else ""
                if f is not None and str(f).strip().isdigit():
                    folder_map[str(f).strip()].append({
                        "farm": "Kibe",
                        "machine": m,
                        "account": str(acc or "").strip(),
                        "name": str(name or "").strip(),
                        "source": "taikhoan_run_safe.xlsx"
                    })
        except Exception as e:
            print(f"[WARN] Lỗi đọc {EXCEL_KIBE_SAFE}: {e}")

    # 2. Admin master taikhoan_dat_v2_updated
    if EXCEL_ADMIN_MASTER.exists():
        try:
            wb = openpyxl.load_workbook(EXCEL_ADMIN_MASTER, data_only=True, read_only=True)
            ws = wb["Tài Khoản"] if "Tài Khoản" in wb.sheetnames else wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 3:
                    continue
                m, f, acc = row[0], row[1], row[2]
                name = row[3] if len(row) > 3 else ""
                if f is not None and str(f).strip().isdigit():
                    folder_map[str(f).strip()].append({
                        "farm": "Admin",
                        "machine": m,
                        "account": str(acc or "").strip(),
                        "name": str(name or "").strip(),
                        "source": "taikhoan_dat_v2_updated .xlsx"
                    })
        except Exception as e:
            print(f"[WARN] Lỗi đọc {EXCEL_ADMIN_MASTER}: {e}")

    return folder_map


def scan_cluster(root_path: Path, cluster_name: str) -> tuple[dict[str, list[str]], int]:
    """Quét thư mục gốc tìm tất cả avatar.jpg và nhóm theo mã băm MD5."""
    if not root_path.exists():
        return {}, 0

    hash_to_folders: dict[str, list[str]] = defaultdict(list)
    total_found = 0
    for p in root_path.glob("*/avatar.jpg"):
        if not p.is_file():
            continue
        folder_num = p.parent.name
        if not folder_num.isdigit():
            continue
        try:
            h = get_md5(p)
            hash_to_folders[h].append(folder_num)
            total_found += 1
        except Exception:
            pass

    return hash_to_folders, total_found


def main() -> int:
    parser = argparse.ArgumentParser(description="Quét toàn bộ avatar bị trùng lặp trên 2 Farm.")
    parser.add_argument("--json-out", default=None, help="Đường dẫn file JSON xuất kết quả chi tiết")
    args = parser.parse_args()

    print("=" * 70)
    print("🔍 BẮT ĐẦU QUÉT TOÀN DIỆN MÃ BĂM AVATAR TRÊN CẢ 2 FARM (KIBE & ADMIN)")
    print("=" * 70)

    mappings = load_account_mappings()

    clusters = [
        ("Kibe Nuôi Nick (D:/TIKTOK-videonuoinick)", ROOT_KIBE_NUOI),
        ("Kibe Video Gốc (D:/video goc)", ROOT_KIBE_GOC),
        ("Admin Video Gốc (D:/video goc may 2)", ROOT_ADMIN_GOC),
    ]

    report_data = {
        "clusters": {},
        "global_duplicates": []
    }

    # Bảng tổng hợp chéo tất cả các root
    global_hash_to_entries: dict[str, list[dict[str, str]]] = defaultdict(list)

    for c_label, c_path in clusters:
        if not c_path.exists():
            print(f"\n[!] Thư mục không tồn tại: {c_path}")
            continue

        h_map, total = scan_cluster(c_path, c_label)
        dups = {h: fl for h, fl in h_map.items() if len(fl) > 1}
        print(f"\n📁 Cụm [{c_label}]:")
        print(f"   • Tổng số avatar tìm thấy: {total}")
        print(f"   • Số nhóm bị trùng lặp nội bộ: {len(dups)}")

        for h, folders in h_map.items():
            for f in folders:
                global_hash_to_entries[h].append({
                    "cluster": c_label,
                    "folder": f,
                    "path": str(c_path / f / "avatar.jpg")
                })

        report_data["clusters"][c_label] = {
            "total_avatars": total,
            "duplicate_groups": len(dups),
        }

    # Tổng hợp các nhóm trùng lặp
    dup_groups = {h: entries for h, entries in global_hash_to_entries.items() if len(entries) > 1}

    print("\n" + "=" * 70)
    print(f"⚠️ KẾT QUẢ ĐỐI SOÁT TRÙNG LẶP TOÀN FARM ({len(dup_groups)} NHÓM TRÙNG):")
    print("=" * 70)

    for idx, (h, entries) in enumerate(sorted(dup_groups.items(), key=lambda x: len(x[1]), reverse=True), 1):
        folders_distinct = sorted(list({e["folder"] for e in entries}), key=lambda x: int(x))
        print(f"\n[{idx}] Nhóm trùng ({len(folders_distinct)} folders distinct, {len(entries)} files) [MD5: {h[:8]}]:")
        print(f"    Folders: {folders_distinct}")
        for f in folders_distinct:
            acc_info = mappings.get(f, [])
            if acc_info:
                for info in acc_info:
                    print(f"      -> Folder {f} | {info['farm']} Máy {info['machine']} | @{info['account']} ({info['name']})")
            else:
                print(f"      -> Folder {f} | (Chưa gán máy hoặc nằm ngoài danh sách workbook active)")

        report_data["global_duplicates"].append({
            "hash": h,
            "folders": folders_distinct,
            "entries_count": len(entries),
        })

    if args.json_out:
        out_p = Path(args.json_out)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        out_p.write_text(json.dumps(report_data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n[✓] Đã xuất báo cáo chi tiết ra file JSON: {out_p}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
