# -*- coding: utf-8 -*-
"""Batch recovery: Write confirmed 2FA secrets from DPAPI journals into taikhoan_dat_v2_updated .xlsx"""
from __future__ import annotations

import os
import sys
import json
import time
from pathlib import Path
from datetime import datetime

# Setup paths
sys.path.insert(0, r"D:\Taadaa\automation-core\src")
sys.path.insert(0, r"D:\Taadaa\tiktok-add-bao-mat-f2a\python_runner")

import openpyxl
from core.journal import JournalStore, PhaseBState, account_hash
from core.workbook import acquire_workbook_lock, resolve_sheet

def run(dry_run: bool = True):
    wb_path = Path(r"D:\OneDrive\TaadaaData\kibe\taikhoan_dat_v2_updated .xlsx")
    j_dir = Path(os.environ["LOCALAPPDATA"]) / "codex_gmail_debug-tiktok-add-bao-mat-f2a" / "journals"
    backup_dir = Path(os.environ["LOCALAPPDATA"]) / "codex_gmail_debug-tiktok-add-bao-mat-f2a" / "workbook-backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    
    if not wb_path.exists():
        print(f"[ERROR] Workbook not found: {wb_path}")
        return 1
        
    store = JournalStore(j_dir)
    
    print(f"Loading workbook: {wb_path}...")
    wb = openpyxl.load_workbook(wb_path, data_only=False, read_only=False)
    ws = resolve_sheet(wb, "Tài Khoản")
    
    # Map headers
    row1 = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    headers = {str(val).upper().strip(): c for c, val in enumerate(row1, start=1) if val}
    
    col_machine = headers.get("MÁY", 1)
    col_id = headers.get("ID", 3)
    col_2fa = headers.get("2FA", 5)
    
    print(f"Header mapping: MÁY={col_machine}, ID={col_id}, 2FA={col_2fa}")
    
    dpapi_files = [f for f in os.listdir(j_dir) if f.endswith(".dpapi")]
    print(f"Found {len(dpapi_files)} journal files in {j_dir}")
    
    targets_to_write = []
    already_done = []
    skipped_mismatch = []
    
    for f in dpapi_files:
        stem = f[:-6]
        try:
            r = store.load(stem)
        except Exception as e:
            continue
            
        row_idx = r.source_row
        m_val = str(ws.cell(row_idx, col_machine).value or "").strip()
        id_val = str(ws.cell(row_idx, col_id).value or "").strip().lstrip("@").lower()
        curr_2fa = str(ws.cell(row_idx, col_2fa).value or "").strip()
        
        expected_m = str(r.machine).strip()
        expected_id = r.username_normalized.strip().lstrip("@").lower()
        
        if m_val != expected_m or id_val != expected_id:
            skipped_mismatch.append((r.machine, row_idx, r.username_normalized, m_val, id_val))
            continue
            
        if curr_2fa:
            already_done.append((stem, r.machine, row_idx, r.username_normalized))
            continue
            
        if r.state == PhaseBState.AUTHENTICATOR_CONFIRMED and r.secret:
            targets_to_write.append((stem, r.machine, row_idx, r.username_normalized, r.secret.strip()))

    print(f"Already populated in Excel: {len(already_done)}")
    print(f"Mismatched identity: {len(skipped_mismatch)}")
    print(f"Targets ready to write: {len(targets_to_write)}")
    
    for stem, m, row_idx, uname, secret in sorted(targets_to_write, key=lambda x: x[2]):
        print(f"  - M{m:2s} (Row {row_idx:3d}): @{uname:20s} | Secret: {secret[:6]}...{secret[-4:]}")
        
    if not targets_to_write:
        print("[INFO] No missing 2FA records to write.")
        wb.close()
        return 0
        
    if dry_run:
        print("\n[DRY RUN] Would write records to Excel. Pass --apply to execute.")
        wb.close()
        return 0
        
    # APPLY WRITE
    print(f"\n[APPLY] Acquiring workbook lock and writing {len(targets_to_write)} records...")
    wb.close() # close read handle
    
    # Make backup copy
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = backup_dir / f"{wb_path.stem}.{stamp}.pre_journal_recovery.xlsx"
    import shutil
    shutil.copy2(wb_path, backup_file)
    print(f"Backup created: {backup_file}")
    
    with acquire_workbook_lock(wb_path, timeout=60.0):
        wb_write = openpyxl.load_workbook(wb_path, data_only=False, read_only=False)
        ws_write = resolve_sheet(wb_write, "Tài Khoản")
        
        written_count = 0
        for stem, m, row_idx, uname, secret in targets_to_write:
            cell = ws_write.cell(row_idx, col_2fa)
            cell.value = secret
            written_count += 1
            
        wb_write.save(wb_path)
        wb_write.close()
        print(f"Saved {written_count} 2FA secrets to {wb_path}")
        
    # Verification
    print("Verifying committed workbook...")
    wb_verify = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    ws_verify = resolve_sheet(wb_verify, "Tài Khoản")
    
    verified_ok = 0
    for stem, m, row_idx, uname, secret in targets_to_write:
        val = str(ws_verify.cell(row_idx, col_2fa).value or "").strip()
        if val == secret:
            verified_ok += 1
            # Purge or transition journal
            try:
                store.purge(stem)
            except Exception:
                pass
        else:
            print(f"[VERIFY FAILED] Row {row_idx}: expected {secret[:6]}..., got {val[:6]}...")
            
    wb_verify.close()
    
    print(f"\n=== RESULT: {verified_ok}/{len(targets_to_write)} records verified & journals purged ===")
    return 0 if verified_ok == len(targets_to_write) else 1

if __name__ == "__main__":
    apply_mode = "--apply" in sys.argv
    sys.exit(run(dry_run=not apply_mode))
