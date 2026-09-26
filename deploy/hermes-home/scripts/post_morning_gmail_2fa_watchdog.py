#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
post_morning_gmail_2fa_watchdog.py
Watchdog cuốn chiếu: Bật 2FA Google Authenticator cho các Profile GPM sau Ca 1 nuôi feed.
Khung giờ: 08:30 - 11:30 (HCM).
Silent Watchdog: Tuyệt đối không in ra stdout nếu không có kết quả hành động.
"""

import os
import sys
import re
import json
import time
import sqlite3
import argparse
import subprocess

from datetime import datetime, timedelta
from pathlib import Path

# Paths
STATE_DIR = Path("D:/Taadaa/runtime/kibe/cron-state")
STATE_FILE = STATE_DIR / "post_morning_gmail_2fa_state.json"
FEED_REPORTED_FILE = STATE_DIR / "feed_session_reported.json"
GPM_DB_PATH = Path(r"C:\Users\Kibe\AppData\Local\Programs\GPMLogin\profile\profile_data.db")
MASTER_EXCEL = Path(r"D:\OneDrive\TaadaaData\kibe\master_gmail_manager.xlsx")
CLEAN_V2_EXCEL = Path(r"D:\OneDrive\TaadaaData\kibe\gmail_clean_v2.xlsx")
NURTURE_STATE_FILE = STATE_DIR / "gpm_gmail_nurture_state.json"
LOCK_FILE = STATE_DIR / "post_morning_gmail_2fa.lock"
LOCK_DIR = Path(r"C:\Users\Kibe\AppData\Local\automation-core\device-locks")

# Import module cron nuôi acc để đọc lịch chạy
sys.path.insert(0, r"D:\Taadaa\tiktok-luot nuoi acc")
try:
    from python_runner.hermes_cron.models import StatePaths, parse_hcm_timestamp
    from python_runner.hermes_cron.manifest import load_active
    from python_runner.hermes_cron.source_config import SourceConfig
    CRON_MODULE_AVAILABLE = True
except Exception:
    CRON_MODULE_AVAILABLE = False

# Add GPM auto script to path for setup_authenticator_for_profile
GPM_SCRIPTS_DIR = Path(r"D:\Taadaa\GPM auto\scripts")
if str(GPM_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(GPM_SCRIPTS_DIR))

try:
    from run_add_2fa_remaining import setup_authenticator_for_profile
except ImportError:
    setup_authenticator_for_profile = None


def log(msg: str):
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sys.stderr.write(f"[{now_str}] {msg}\n")


def _profile_has_google_session(profile) -> bool:
    """Return whether the GPM profile has Google session cookies.

    Candidates carry the verification result so the action consumer does not
    confuse a missing field with a missing session.  The fallback keeps this
    helper safe for callers/tests that provide only the profile path.
    """
    if not isinstance(profile, dict):
        return False
    if profile.get("google_session_verified") is True or profile.get("session_verified") is True:
        return True
    if profile.get("google_session_verified") is False or profile.get("session_verified") is False:
        return False

    profile_path = profile.get("profile_path")
    if not profile_path:
        return False
    p_folder = Path(profile_path)
    if not p_folder.is_absolute():
        p_folder = GPM_DB_PATH.parent / p_folder
    for cookie_path in (p_folder / "Default" / "Network" / "Cookies", p_folder / "Default" / "Cookies"):
        if not cookie_path.exists():
            continue
        try:
            conn = sqlite3.connect(str(cookie_path))
            cur = conn.cursor()
            cur.execute(
                "SELECT count(*) FROM cookies WHERE host_key LIKE '%google.com' "
                "AND name IN ('SID', 'SSID', 'HSID', 'SAPISID')"
            )
            has_session = cur.fetchone()[0] >= 2
            conn.close()
            if has_session:
                return True
        except Exception:
            continue
    return False


class ProcessLock:
    """Small cross-process lock for the watchdog's single-writer critical section."""

    def __init__(self, lock_path: str | Path):
        self.lock_path = str(lock_path)
        self._file = None
        self.acquired = False

    def acquire(self) -> bool:
        if self.acquired:
            return False
        handle = None
        try:
            Path(self.lock_path).parent.mkdir(parents=True, exist_ok=True)
            handle = open(self.lock_path, "a+b")
            handle.seek(0)
            handle.write(b"\0")
            handle.flush()
            handle.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            self._file = handle
            self.acquired = True
            return True
        except (OSError, IOError):
            if handle:
                handle.close()
            return False

    def release(self):
        if not self.acquired or not self._file:
            return
        try:
            if os.name == "nt":
                import msvcrt
                self._file.seek(0)
                msvcrt.locking(self._file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(self._file.fileno(), fcntl.LOCK_UN)
        except OSError:
            pass
        finally:
            self._file.close()
            self._file = None
            self.acquired = False


def is_within_time_window(start_hm=(8, 30), end_hm=(11, 30)) -> bool:
    now = datetime.now()
    cur_hm = (now.hour, now.minute)
    return start_hm <= cur_hm <= end_hm


def is_feed_ca1_finished() -> bool:
    if not FEED_REPORTED_FILE.exists():
        return False
    try:
        with open(FEED_REPORTED_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        today = datetime.now().strftime("%Y-%m-%d")
        keys = [f"{today}_ca1_phien2", f"{today}_ca1", f"{today}_ca1_phien3"]
        if isinstance(data, dict):
            reported = set(data.get("reported_sessions", []))
            return any(k in reported for k in keys) or any(k in data and data[k] for k in keys)
        elif isinstance(data, list):
            return any(isinstance(item, str) and any(k in item for k in keys) for item in data)
        return False
    except Exception as e:
        log(f"Lỗi kiểm tra feed session reported: {e}")
        return False


def is_feed_runner_active() -> bool:
    try:
        if os.name == "nt":
            ps_cmd = 'powershell -NoProfile -Command "Get-CimInstance Win32_Process -Filter \\"Name LIKE \'python%\'\\" | Select-Object -ExpandProperty CommandLine"'
            ps_out = subprocess.check_output(ps_cmd, shell=True, text=True, errors="ignore", timeout=15)
            keywords = ["feed_session_runner", "run_feed_session", "feed_consumer", "multi_machine_feed_session"]
            for line in ps_out.splitlines():
                if any(kw in line.lower() for kw in keywords):
                    return True
        else:
            res = subprocess.run(["pgrep", "-f", "feed_session_runner|run_feed_session|feed_consumer"],
                                 capture_output=True, text=True, timeout=5)
            if res.returncode == 0:
                return True
        return False
    except Exception as e:
        log(f"Lỗi kiểm tra feed runner: {e}")
        return False


def get_active_locks() -> set:
    active = set()
    if LOCK_DIR.exists():
        for f in LOCK_DIR.glob("*.json"):
            try:
                d = json.loads(f.read_text(encoding="utf-8"))
                if d.get("status") in ["active", "running", "queued", "blocked"]:
                    m = d.get("machine")
                    if m is not None:
                        active.add(int(m))
            except Exception:
                pass
    return active


def get_upcoming_feed_machines(buffer_minutes=60) -> set:
    if not CRON_MODULE_AVAILABLE:
        return set()
    try:
        root = Path("D:/Taadaa/runtime/kibe/cron-state")
        source_file = Path("D:/Taadaa/runtime/kibe/cron-source/hermes_cron_source_config.json")
        if not root.exists() or not source_file.exists():
            return set()
        source = SourceConfig.from_json(source_file)
        now_dt = datetime.now().astimezone(parse_hcm_timestamp("2026-08-18T00:00:00+07:00").tzinfo)
        day_str = now_dt.strftime("%Y-%m-%d")
        active = load_active(StatePaths(root, root), day_str, source)
        busy_window_end = now_dt + timedelta(minutes=buffer_minutes)
        busy_machines = set()
        for entry in active.payload.get("entries", []):
            m = entry.get("machine")
            if m is None:
                continue
            s_start = parse_hcm_timestamp(entry["slot_time"])
            s_end = parse_hcm_timestamp(entry["slot_end"])
            if (s_start <= now_dt < s_end) or (now_dt <= s_start < busy_window_end):
                busy_machines.add(int(m))
        return busy_machines
    except Exception as e:
        log(f"Lỗi khi đọc manifest cron nuôi acc: {e}")
        return set()


def get_gpm_candidates(limit: int = 5) -> list:
    """Lấy danh sách profile GPM cần bật 2FA (có email @gmail.com, có password, chưa có 2FA)."""
    if not GPM_DB_PATH.exists() or not MASTER_EXCEL.exists():
        return []

    import openpyxl

    master_pwd = {}
    master_2fa = {}
    try:
        wb_m = openpyxl.load_workbook(MASTER_EXCEL, read_only=True, data_only=True)
        if "Master_All" in wb_m.sheetnames:
            ws = wb_m["Master_All"]
            for r in list(ws.iter_rows(values_only=True))[1:]:
                em = str(r[1] or "").strip().lower()
                if em:
                    master_pwd[em] = str(r[2] or "").strip()
                    if len(r) > 4 and r[4]:
                        master_2fa[em] = str(r[4] or "").strip()
        wb_m.close()
    except Exception as e:
        log(f"Lỗi đọc MASTER_EXCEL: {e}")

    clean_pwd = {}
    clean_2fa = {}
    if CLEAN_V2_EXCEL.exists():
        try:
            wb_c = openpyxl.load_workbook(CLEAN_V2_EXCEL, read_only=True, data_only=True)
            ws_c = wb_c.active
            for r in list(ws_c.iter_rows(values_only=True))[1:]:
                em = str(r[1] or "").strip().lower()
                if em:
                    clean_pwd[em] = str(r[2] or "").strip()
                    if len(r) > 3 and r[3]:
                        clean_2fa[em] = str(r[3] or "").strip()
            wb_c.close()
        except Exception as e:
            log(f"Lỗi đọc CLEAN_V2_EXCEL: {e}")

    candidates = []
    try:
        conn = sqlite3.connect(GPM_DB_PATH)
        c = conn.cursor()
        c.execute("SELECT Id, Name, ProfilePath, GroupId FROM Profiles ORDER BY Id ASC")
        rows = c.fetchall()
        conn.close()

        for pid, name, ppath, gid in rows:
            if not name or name.startswith("AMZ_") or "deleted" in name.lower():
                continue
            if gid != 10:
                continue
            em_match = re.search(r"([a-zA-Z0-9_.+-]+@gmail\.com)", name, re.IGNORECASE)
            if not em_match:
                continue
            email = em_match.group(1).lower()
            if "khoale" in email:
                continue

            sec = master_2fa.get(email) or clean_2fa.get(email) or ""
            if sec and sec not in ("", "None", "null"):
                continue

            pwd = master_pwd.get(email) or clean_pwd.get(email) or ""
            if not pwd:
                continue

            # SOAKING GATE: Bắt buộc đã được nuôi ít nhất 1 phiên thành công trên GPM trước khi bật 2FA
            if NURTURE_STATE_FILE.exists():
                try:
                    if "nurture_cache" not in locals():
                        nurture_cache = json.loads(NURTURE_STATE_FILE.read_text(encoding="utf-8"))
                    n_info = nurture_cache.get(email, {})
                    if not n_info.get("last_nurtured") or n_info.get("status") != "success":
                        log(f"[2FA-SOAK-GATE] Bỏ qua {email}: Đang ngâm trên GPM (chưa có phiên nuôi thành công)")
                        continue
                except Exception:
                    pass

            m_match = re.match(r"^(\d+)", name) or re.search(r"M(\d+)", name, re.IGNORECASE)
            machine_num = int(m_match.group(1)) if m_match else 0

            candidates.append({
                "id": pid,
                "name": name,
                "email": email,
                "machine": machine_num,
                "pwd": pwd,
                "2fa_secret": "",
                "group_id": gid,
                "profile_path": ppath,
                "google_session_verified": _profile_has_google_session({"profile_path": ppath}),
            })
            if len(candidates) >= limit:
                break
    except Exception as e:
        log(f"Lỗi truy vấn GPM database: {e}")

    return candidates


def save_state_results(success_list: list, fail_list: list, pending_list: list | None = None):
    try:
        STATE_DIR.mkdir(parents=True, exist_ok=True)
        today = datetime.now().strftime("%Y-%m-%d")
        data = {
            "last_run_date": today,
            "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "success": success_list,
            "failed": fail_list,
            "pending": pending_list or []
        }
        with open(STATE_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
    except Exception as e:
        log(f"Lỗi ghi state: {e}")


def main():
    parser = argparse.ArgumentParser(description="Watchdog bật 2FA cho profile GPM sau Ca 1.")
    parser.add_argument("--dry-run", action="store_true", help="Chỉ quét candidates, không thực thi")
    parser.add_argument("--force", action="store_true", help="Bỏ qua kiểm tra giờ và feed session")
    parser.add_argument("--batch-size", type=int, default=5, help="Số profile xử lý tối đa mỗi tick")
    args = parser.parse_args()

    if not args.force and not is_within_time_window((8, 30), (11, 30)):
        return 0

    if not args.force and not is_feed_ca1_finished():
        return 0

    if not args.force and is_feed_runner_active():
        log("Feed runner đang chạy, hoãn bật 2FA GPM.")
        return 0

    candidates = get_gpm_candidates(limit=args.batch_size)
    if not candidates:
        return 0

    active_locks = get_active_locks() if not args.force else set()
    busy_feed = get_upcoming_feed_machines(buffer_minutes=60) if not args.force else set()

    eligible_candidates = []
    for c in candidates:
        m = c.get("machine", 0)
        if not args.force and m > 0:
            if m in active_locks:
                log(f"[LOCK-SKIP] Bỏ qua {c['email']} (Máy {m}): máy đang có lock ({active_locks})")
                continue
            if m in busy_feed:
                log(f"[FEED-SKIP] Bỏ qua {c['email']} (Máy {m}): máy đang hoặc sắp chạy feed ({busy_feed})")
                continue
        eligible_candidates.append(c)

    if not eligible_candidates:
        return 0

    if args.dry_run:
        log(f"[DRY-RUN] Tìm thấy {len(eligible_candidates)} profile GPM cần bật 2FA:")
        for c in eligible_candidates:
            log(f"  - Máy {c['machine']:02d} | {c['name']} ({c['email']})")
        return 0

    if not setup_authenticator_for_profile:
        log("Lỗi: Không thể import setup_authenticator_for_profile từ GPM auto")
        return 1

    success_list = []
    fail_list = []
    pending_list = []

    lock = ProcessLock(LOCK_FILE)
    if not lock.acquire():
        log("Tiến trình 2FA đang chạy ngầm, im lặng thoát.")
        return 0
    try:
        for c in eligible_candidates:
            try:
                log(f"Bắt đầu xử lý 2FA GPM: {c['name']} ({c['email']})...")
                if not _profile_has_google_session(c):
                    log(f"[GPM-2FA-SKIP] {c['email']}: PENDING_GPM_LOGIN_NO_SESSION")
                    pending_list.append(f"{c['email']}: PENDING_GPM_LOGIN_NO_SESSION")
                    continue
                res = setup_authenticator_for_profile(c)
                if res.get("status") == "SUCCESS" and res.get("secret_key"):
                    success_list.append(f"{c['email']} ({res['secret_key']})")
                elif res.get("status") == "ALREADY_ACTIVE":
                    success_list.append(f"{c['email']} (ALREADY_ACTIVE)")
                else:
                    fail_list.append(f"{c['email']}: {res.get('status')} - {res.get('details')}")
            except Exception as e:
                log(f"Exception khi xử lý {c['email']}: {e}")
                fail_list.append(f"{c['email']}: {str(e)}")
    finally:
        lock.release()

    if success_list or fail_list or pending_list:
        save_state_results(success_list, fail_list, pending_list)
        report = [
            "### [BÁO CÁO 2FA GMAIL QUA GPM PROFILE]",
            f"- Đã xử lý: {len(eligible_candidates)} profile",
            f"- Thành công ({len(success_list)}):",
        ]
        for s in success_list:
            report.append(f"  + {s}")
        if fail_list:
            report.append(f"- Cần lưu ý ({len(fail_list)}):")
            for f_item in fail_list:
                report.append(f"  + {f_item}")
        if pending_list:
            report.append(f"- Đang chờ đăng nhập GPM ({len(pending_list)}):")
            for p_item in pending_list:
                report.append(f"  + {p_item}")
        print("\n".join(report))

    return 0


if __name__ == "__main__":
    sys.exit(main())
