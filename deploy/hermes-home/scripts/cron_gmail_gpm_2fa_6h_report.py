#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Báo cáo định kỳ mỗi 6h tình trạng 2FA Gmail qua GPMLogin Profile.
Deliver: telegram:-5373649734 (Farm Alerts)

Quy ước I/O: stdout CHỈ chứa nội dung báo cáo (được cron gửi lên Telegram);
mọi log/telemetry đi qua stderr. Không có sự kiện trong 6h -> stdout rỗng.
"""

import json
import logging
import re
import sqlite3
import sys
import zipfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

GPM_DB_PATH = Path(r"C:\Users\Kibe\AppData\Local\Programs\GPMLogin\profile\profile_data.db")
MASTER_EXCEL = Path(r"D:\OneDrive\TaadaaData\kibe\master_gmail_manager.xlsx")
CLEAN_V2_EXCEL = Path(r"D:\OneDrive\TaadaaData\kibe\gmail_clean_v2.xlsx")
NURTURE_STATE_FILE = Path(r"D:\Taadaa\runtime\kibe\cron-state\gpm_gmail_nurture_state.json")
STATE_FILE = Path(r"D:\Taadaa\runtime\kibe\cron-state\post_morning_gmail_2fa_state.json")

GPM_GROUP_ID = 10
REPORT_WINDOW_HOURS = 6
EMPTY_SECRET_VALUES = ("", "None", "null")
GMAIL_RE = re.compile(r"([a-zA-Z0-9_.+-]+@gmail\.com)", re.IGNORECASE)

logger = logging.getLogger("cron_gmail_gpm_2fa_6h_report")


def setup_logging() -> None:
    """Cấu hình log ra stderr để không làm bẩn stdout của báo cáo."""
    if logger.handlers:
        return
    handler = logging.StreamHandler(sys.stderr)
    handler.setFormatter(logging.Formatter("[%(asctime)s] %(levelname)s %(message)s"))
    logger.addHandler(handler)
    logger.setLevel(logging.INFO)


def emit_metric(event: str, **fields: Any) -> None:
    payload = {"event": event, **fields}
    logger.info("[TELEMETRY_METRIC] %s", json.dumps(payload, ensure_ascii=False))


def _excel_errors() -> Tuple[type, ...]:
    from openpyxl.utils.exceptions import InvalidFileException

    return (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException)


def _read_secret_column(
    path: Path, email_col: int, secret_col: int, sheet_name: Optional[str], label: str, failed: List[str]
) -> Dict[str, str]:
    """Đọc map email -> secret 2FA từ một file Excel; lỗi được ghi log và đánh dấu nguồn hỏng."""
    secrets: Dict[str, str] = {}
    if not path.exists():
        logger.warning("%s không tồn tại: %s", label, path)
        return secrets
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name is None:
                ws = wb.active
            elif sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                logger.warning("%s thiếu sheet %s", label, sheet_name)
                return secrets
            for row in list(ws.iter_rows(values_only=True))[1:]:
                if len(row) <= secret_col:
                    continue
                email = str(row[email_col] or "").strip().lower()
                if email and row[secret_col]:
                    secrets[email] = str(row[secret_col] or "").strip()
        finally:
            wb.close()
    except _excel_errors() as exc:
        logger.error("Lỗi đọc %s (%s): %s: %s", label, path, type(exc).__name__, exc)
        failed.append(label)
    return secrets


def load_master_2fa(failed: List[str]) -> Dict[str, str]:
    return _read_secret_column(MASTER_EXCEL, 1, 4, "Master_All", "MASTER_EXCEL", failed)


def load_clean_2fa(failed: List[str]) -> Dict[str, str]:
    return _read_secret_column(CLEAN_V2_EXCEL, 1, 3, None, "CLEAN_V2_EXCEL", failed)


def load_json_file(path: Path, label: str, failed: List[str]) -> Dict[str, Any]:
    """Đọc JSON dạng dict; file thiếu -> {} (không phải lỗi), file hỏng/sai kiểu -> {} + log lỗi."""
    if not path.exists():
        logger.warning("%s không tồn tại: %s", label, path)
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        logger.error("Lỗi đọc %s (%s): %s: %s", label, path, type(exc).__name__, exc)
        failed.append(label)
        return {}
    if not isinstance(data, dict):
        logger.error("%s không phải JSON object (got %s)", label, type(data).__name__)
        failed.append(label)
        return {}
    return data


def load_group_profile_names(failed: List[str]) -> Optional[List[str]]:
    """Trả về danh sách tên profile của group 10; None nếu DB không đọc được."""
    if not GPM_DB_PATH.exists():
        logger.warning("GPM_DB không tồn tại: %s", GPM_DB_PATH)
        return None
    try:
        conn = sqlite3.connect(GPM_DB_PATH)
        try:
            rows = conn.execute(
                "SELECT Id, Name FROM Profiles WHERE GroupId = ? ORDER BY Id ASC", (GPM_GROUP_ID,)
            ).fetchall()
        finally:
            conn.close()
    except sqlite3.Error as exc:
        logger.error("Lỗi đọc GPM_DB (%s): %s: %s", GPM_DB_PATH, type(exc).__name__, exc)
        failed.append("GPM_DB")
        return None
    return [name or "" for _pid, name in rows]


def extract_email(profile_name: str) -> Optional[str]:
    match = GMAIL_RE.search(profile_name or "")
    return match.group(1).lower() if match else None


def classify_profiles(
    names: List[str],
    master_2fa: Dict[str, str],
    clean_2fa: Dict[str, str],
    nurture_cache: Dict[str, Any],
) -> Dict[str, int]:
    """Phân loại profile: đã bật 2FA / đang ngâm / sẵn sàng bật 2FA. Profile không có Gmail hoặc 'khoale' bị bỏ qua."""
    stats = {"total": len(names), "already_2fa": 0, "soaking": 0, "ready": 0, "skipped": 0}
    for name in names:
        email = extract_email(name)
        if not email or "khoale" in email:
            stats["skipped"] += 1
            continue
        secret = master_2fa.get(email) or clean_2fa.get(email) or ""
        if secret not in EMPTY_SECRET_VALUES:
            stats["already_2fa"] += 1
            continue
        info = nurture_cache.get(email)
        if not isinstance(info, dict) or not info.get("last_nurtured") or info.get("status") != "success":
            stats["soaking"] += 1
            continue
        stats["ready"] += 1
    return stats


def get_needs_login_profiles(names: List[str], nurture_cache: Dict[str, Any]) -> List[str]:
    """Trích xuất danh sách profile trong group đang bị văng cookie (NEEDS_LOGIN)."""
    dropped = []
    for name in names:
        email = extract_email(name)
        if not email or "khoale" in email:
            continue
        info = nurture_cache.get(email)
        if isinstance(info, dict) and info.get("status") == "NEEDS_LOGIN":
            dropped.append(email)
    return dropped


def parse_event_time(raw: str) -> datetime:
    """Parse ISO timestamp thành datetime naive theo giờ local (hậu tố 'Z' coi như giờ local, như cũ)."""
    parsed = datetime.fromisoformat(raw.replace("Z", ""))
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone().replace(tzinfo=None)
    return parsed


def filter_recent_events(events: Any, since: datetime) -> List[Dict[str, Any]]:
    if not isinstance(events, list):
        logger.error("STATE_FILE.events không phải list (got %s)", type(events).__name__)
        return []
    recent: List[Dict[str, Any]] = []
    for ev in events:
        if not isinstance(ev, dict):
            logger.warning("Bỏ qua event không phải dict: %r", ev)
            continue
        at_str = ev.get("at", "")
        if not at_str:
            continue
        try:
            if parse_event_time(str(at_str)) >= since:
                recent.append(ev)
        except ValueError as exc:
            logger.warning("Lỗi parse timestamp event %r: %s", at_str, exc)
    return recent


def format_event(ev: Dict[str, Any]) -> str:
    machine = ev.get("machine", 0)
    try:
        m_str = f"M{int(machine):02d}" if machine else "M??"
    except (TypeError, ValueError):
        m_str = "M??"
    email = ev.get("email", "")
    status = ev.get("status", "")
    if status == "SUCCESS":
        secret = str(ev.get("secret_key") or "")
        short = f" ({secret[:8]}...{secret[-4:]})" if len(secret) > 12 else (f" ({secret})" if secret else "")
        return f"  + {m_str} | {email}: SUCCESS{short}"
    if status == "ALREADY_ACTIVE":
        return f"  + {m_str} | {email}: ALREADY_ACTIVE"
    return f"  + {m_str} | {email}: {status} - {str(ev.get('details') or '')[:60]}"


def render_report(stats: Dict[str, int], events: List[Dict[str, Any]], failed: List[str], needs_login: Optional[List[str]] = None) -> str:
    total = stats["total"]
    pct = round(stats["already_2fa"] / total * 100, 1) if total > 0 else 0
    lines = [
        "📊 [BÁO CÁO 6H: 2FA GMAIL QUA GPM PROFILE]",
        f"• Tổng số profile GPM (Group {GPM_GROUP_ID}): {total}",
        f"• Đã bật 2FA thành công: {stats['already_2fa']}/{total} ({pct}%)",
        f"• Đang ngâm (chờ phiên nuôi): {stats['soaking']}",
        f"• Sẵn sàng bật 2FA: {stats['ready']}",
    ]
    if needs_login:
        lines.append(f"⚠️ Acc văng cookie (NEEDS_LOGIN): {len(needs_login)} profile")
        for em in needs_login[:5]:
            lines.append(f"  - {em}")
        if len(needs_login) > 5:
            lines.append(f"  ... và {len(needs_login) - 5} acc khác")
    if failed:
        lines.append(f"⚠️ Nguồn dữ liệu lỗi (số liệu có thể thiếu): {', '.join(failed)}")
    lines.append("")
    lines.append(f"⚡ Tiến độ xử lý trong {REPORT_WINDOW_HOURS}h vừa qua: {len(events)} sự kiện")
    lines.extend(format_event(ev) for ev in events)
    return "\n".join(lines)


def main() -> int:
    setup_logging()
    started = datetime.now()
    failed: List[str] = []

    master_2fa = load_master_2fa(failed)
    clean_2fa = load_clean_2fa(failed)
    nurture_cache = load_json_file(NURTURE_STATE_FILE, "NURTURE_STATE_FILE", failed)

    names = load_group_profile_names(failed)
    stats = classify_profiles(names or [], master_2fa, clean_2fa, nurture_cache)
    needs_login = get_needs_login_profiles(names or [], nurture_cache)

    state = load_json_file(STATE_FILE, "STATE_FILE", failed)
    recent_events = filter_recent_events(
        state.get("events", []), datetime.now() - timedelta(hours=REPORT_WINDOW_HOURS)
    )
    silent = not recent_events

    emit_metric(
        "gmail_2fa_6h_report",
        group=GPM_GROUP_ID,
        total=stats["total"],
        already_2fa=stats["already_2fa"],
        soaking=stats["soaking"],
        ready=stats["ready"],
        needs_login=len(needs_login),
        skipped=stats["skipped"],
        recent_events=len(recent_events),
        silent=silent,
        failed_sources=failed,
        duration_s=round((datetime.now() - started).total_seconds(), 3),
    )

    # SILENT WATCHDOG: không có sự kiện mới trong 6h -> im lặng tuyệt đối (chống spam Telegram)
    if silent:
        return 0

    print(render_report(stats, recent_events, failed, needs_login))
    return 0


if __name__ == "__main__":
    sys.exit(main())
