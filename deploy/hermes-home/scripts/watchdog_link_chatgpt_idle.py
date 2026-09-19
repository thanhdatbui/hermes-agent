#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
watchdog_link_chatgpt_idle.py
Watchdog cuốn chiếu: Tự động chạy liên kết ChatGPT trên máy S7 khi máy rảnh giữa các ca nuôi TikTok,
đồng thời tự động gỡ bỏ tài khoản Gmail DIE khỏi thiết bị S7 khi phát hiện máy rảnh.
"""
from __future__ import annotations

import os
import sys
import json
import time
import logging
import subprocess
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S")
logger = logging.getLogger("watchdog_link_chatgpt_idle")

HCMC = ZoneInfo("Asia/Ho_Chi_Minh")
ADB_EXE = r"C:\Program Files (x86)\xiaowei\tools\adb.exe"
if not os.path.exists(ADB_EXE):
    ADB_EXE = r"C:\Users\Kibe\.GemPhoneFarm\app\adb-tool\adb.exe"

STATE_DIR = Path(r"D:\Taadaa\runtime\kibe\cron-state")
STATE_FILE = STATE_DIR / "chatgpt_link_backlog_state.json"
LOCK_DIR = Path(r"C:\Users\Kibe\AppData\Local\automation-core\device-locks")
MANIFEST_DIR = Path(r"D:\Taadaa\runtime\kibe\cron-state\manifests")
DIE_TONG_FILE = Path(r"D:\OneDrive\TaadaaData\kibe\gmail_die_tong.txt")
MASTER_XLSX = Path(r"D:\OneDrive\TaadaaData\kibe\master_gmail_manager.xlsx")

def get_live_targets() -> list[dict]:
    """Quét danh sách các tài khoản Gmail LIVE chưa có ChatGPT từ Master Excel."""
    targets = []
    if not MASTER_XLSX.exists():
        return targets
    try:
        import openpyxl, re
        wb = openpyxl.load_workbook(MASTER_XLSX, data_only=True)
        if "Kibe_Farm_S7" not in wb.sheetnames:
            return targets
        ws = wb["Kibe_Farm_S7"]
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if not r or len(r) < 10:
                continue
            email = str(r[1] or "").strip().lower()
            if not email or "@" not in email:
                continue
            status = str(r[6] or "").strip().upper()
            if status != "LIVE":
                continue
            note = str(r[13] or "").strip().upper()
            if "CHATGPT" in note:
                continue

            m_val = str(r[7] or "")
            m_num = re.search(r"(\d+)", m_val)
            mid = int(m_num.group(1)) if m_num else 0
            serial = str(r[9] or "").strip()
            pwd = str(r[2] or "").strip()
            if mid and serial:
                targets.append({
                    "stt": mid,
                    "serial": serial,
                    "email": email,
                    "pwd": pwd,
                    "dob": "2000-01-01",
                })
        wb.close()
    except Exception as exc:
        sys.stderr.write(f"[TARGETS-ERR] Lỗi đọc danh sách từ Excel: {exc}\n")
    return targets

# 2. Danh sách 7 máy có Gmail DIE cần gỡ khỏi máy khi rảnh:
DIE_TARGETS = [
    {"stt": 5,  "serial": "9885e64b4a434a3037", "email": "tongdong99auog916@gmail.com"},
    {"stt": 11, "serial": "988633474f4b514436", "email": "tuxuandbn6837l@gmail.com"},
    {"stt": 12, "serial": "9886783459534f4a59", "email": "ungvinh.pro.qyjl43@gmail.com"},
    {"stt": 27, "serial": "ce031823912ae0d20c", "email": "lucminhvu2000161@gmail.com"},
    {"stt": 31, "serial": "ce0416041bdb271305", "email": "su.cong.binh.051144@gmail.com"},
    {"stt": 52, "serial": "ce0418243a6250430c", "email": "ha.tieu.2402tfi53@gmail.com"},
    {"stt": 63, "serial": "ce091609dc2dc92804", "email": "lieu.giang.brcp393@gmail.com"},
]


def load_state() -> dict:
    if not STATE_FILE.exists():
        return {"completed_chatgpt": {}, "cleaned_die": {}}
    try:
        data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
        data.setdefault("completed_chatgpt", {})
        data.setdefault("cleaned_die", {})
        return data
    except Exception:
        return {"completed_chatgpt": {}, "cleaned_die": {}}


def save_state(state: dict):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    STATE_FILE.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")


def mark_chatgpt_ready_excel(email: str):
    """Cập nhật cờ CHATGPT_READY vào cột Ghi Chú (cột 14) của master_gmail_manager.xlsx an toàn."""
    if not MASTER_XLSX.exists():
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(MASTER_XLSX)
        if "Kibe_Farm_S7" in wb.sheetnames:
            ws = wb["Kibe_Farm_S7"]
            for row in range(2, ws.max_row + 1):
                cell_em = ws.cell(row=row, column=2).value
                if cell_em and str(cell_em).strip().lower() == email.strip().lower():
                    cur_val = str(ws.cell(row=row, column=14).value or "").strip()
                    if "CHATGPT_READY" not in cur_val.upper():
                        new_val = f"{cur_val} | CHATGPT_READY".strip(" |") if cur_val else "CHATGPT_READY"
                        ws.cell(row=row, column=14, value=new_val)
                    break
            wb.save(MASTER_XLSX)
            wb.close()
    except Exception as exc:
        sys.stderr.write(f"[EXCEL-LOCK] Không thể cập nhật master_gmail_manager.xlsx: {exc}\n")


def is_machine_locked(machine_id: int, serial: str) -> bool:
    if not LOCK_DIR.exists():
        return False
    for f in LOCK_DIR.glob("*.json"):
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
            if d.get("status") in ["active", "running", "queued", "blocked"]:
                if d.get("machine") == machine_id or d.get("serial") == serial:
                    return True
        except Exception:
            pass
    return False


def is_machine_in_feed(serial: str) -> bool:
    try:
        res = subprocess.run(
            [ADB_EXE, "-s", serial, "shell", "dumpsys", "window", "windows"],
            capture_output=True, text=True, timeout=6
        )
        if "com.ss.android.ugc.trill" in res.stdout:
            return True
    except Exception:
        pass
    return False


def get_next_feed_slot_distance_minutes(machine_id: int, now_dt: datetime) -> float:
    today_str = now_dt.strftime("%Y-%m-%d")
    manifest_folder = MANIFEST_DIR / today_str
    if not manifest_folder.exists():
        return 999.0
    manifest_files = list(manifest_folder.glob("assignment-v1-*.json"))
    if not manifest_files:
        return 999.0
    try:
        data = json.loads(manifest_files[0].read_text(encoding="utf-8"))
        min_distance = 999.0
        for e in data.get("entries", []):
            if e.get("machine") == machine_id:
                st_str = e.get("slot_time")
                se_str = e.get("slot_end")
                if not st_str or not se_str:
                    continue
                st_dt = datetime.fromisoformat(st_str)
                se_dt = datetime.fromisoformat(se_str)
                if st_dt <= now_dt < se_dt:
                    return 0.0
                if st_dt > now_dt:
                    dist = (st_dt - now_dt).total_seconds() / 60.0
                    if dist < min_distance:
                        min_distance = dist
        return min_distance
    except Exception:
        return 999.0


def check_and_run_one(dry_run: bool = False) -> dict:
    now = datetime.now(HCMC)
    state = load_state()
    completed_cg = state.get("completed_chatgpt", {})
    cleaned_die = state.get("cleaned_die", {})

    telemetry = {
        "scanned_targets": 0,
        "skipped_locked": 0,
        "skipped_feed": 0,
        "skipped_distance": 0,
        "skipped_die": 0,
        "attempted": 0,
        "success": 0,
        "failed": 0,
        "dry_run": dry_run
    }

    live_targets = get_live_targets()
    pending_cg = [t for t in live_targets if t["email"] not in completed_cg]
    pending_die = [t for t in DIE_TARGETS if t["email"] not in cleaned_die]

    if not pending_cg and not pending_die:
        return telemetry

    # Ưu tiên 1: Chạy liên kết ChatGPT cho máy LIVE rảnh
    for t in pending_cg:
        stt = t["stt"]
        serial = t["serial"]
        email = t["email"]
        telemetry["scanned_targets"] += 1

        if is_machine_locked(stt, serial):
            telemetry["skipped_locked"] += 1
            continue
        dist = get_next_feed_slot_distance_minutes(stt, now)
        if dist < 25.0:
            telemetry["skipped_distance"] += 1
            continue
        if is_machine_in_feed(serial):
            telemetry["skipped_feed"] += 1
            continue

        if dry_run:
            telemetry["attempted"] += 1
            logger.info(f"[DRY-RUN] Máy {stt:02d} ({serial}) sẵn sàng liên kết ChatGPT: {email}")
            return telemetry

        # Kiểm tra Live Gmail trước khi chiếm máy
        sys.path.insert(0, r"D:\Taadaa\tools")
        try:
            from check_gmail_live_fast import check_gmail_is_live
            if not check_gmail_is_live(email):
                telemetry["skipped_die"] += 1
                logger.warning(f"Máy {stt:02d}: Gmail {email} đã DIE trên checkmail.live -> Bỏ qua ChatGPT")
                continue
        except Exception:
            pass

        # Tiến hành liên kết ChatGPT
        telemetry["attempted"] += 1
        logger.info(f"Máy {stt:02d} ({serial}) RẢNH an toàn -> Liên kết ChatGPT: {email}...")
        sys.path.insert(0, r"D:\Taadaa\register gmail\scripts")
        try:
            from hook_chatgpt_register import register_chatgpt_on_device
            res = register_chatgpt_on_device(
                device_id=serial,
                email=email,
                password=t["pwd"],
                dob=t["dob"],
                timeout=180
            )
            if isinstance(res, dict) and res.get("success"):
                telemetry["success"] += 1
                completed_cg[email] = {
                    "machine": stt,
                    "serial": serial,
                    "linked_at": datetime.now(HCMC).isoformat(),
                    "message": res.get("message")
                }
                state["completed_chatgpt"] = completed_cg
                save_state(state)
                mark_chatgpt_ready_excel(email)
                logger.info(f"[CHATGPT-LINKED] Máy {stt:02d} ({email}): {res.get('message')}")
            else:
                telemetry["failed"] += 1
                err_msg = res.get("message") if isinstance(res, dict) else str(res)
                logger.error(f"[CHATGPT-FAIL] Máy {stt:02d} ({email}): {err_msg}")
        except Exception as exc:
            telemetry["failed"] += 1
            logger.error(f"[CHATGPT-EXC] Máy {stt:02d} ({email}): {exc}")
        return telemetry

    # Ưu tiên 2: Dọn dẹp tài khoản DIE trên máy S7 rảnh
    for t in pending_die:
        stt = t["stt"]
        serial = t["serial"]
        email = t["email"]
        telemetry["scanned_targets"] += 1

        if is_machine_locked(stt, serial):
            telemetry["skipped_locked"] += 1
            continue
        dist = get_next_feed_slot_distance_minutes(stt, now)
        if dist < 15.0:  # Gỡ account chỉ mất ~10-15s, cần đệm 15 phút
            telemetry["skipped_distance"] += 1
            continue
        if is_machine_in_feed(serial):
            telemetry["skipped_feed"] += 1
            continue

        if dry_run:
            telemetry["attempted"] += 1
            logger.info(f"[DRY-RUN] Máy {stt:02d} ({serial}) sẵn sàng gỡ Gmail DIE: {email}")
            return telemetry

        telemetry["attempted"] += 1
        logger.info(f"Máy {stt:02d} ({serial}) RẢNH an toàn -> Gỡ Gmail DIE: {email}...")
        sys.path.insert(0, r"D:\Taadaa\tools")
        try:
            from remove_device_google_account import remove_device_account_fast
            ok = remove_device_account_fast(serial, email)
            if ok:
                telemetry["success"] += 1
            else:
                telemetry["failed"] += 1
            cleaned_die[email] = {
                "machine": stt,
                "serial": serial,
                "cleaned_at": datetime.now(HCMC).isoformat(),
                "success": bool(ok)
            }
            state["cleaned_die"] = cleaned_die
            save_state(state)
            logger.info(f"[DIE-CLEANED] Máy {stt:02d}: Đã gỡ tài khoản DIE {email} khỏi thiết bị S7")
        except Exception as exc:
            telemetry["failed"] += 1
            logger.error(f"[DIE-CLEAN-EXC] Máy {stt:02d} ({email}): {exc}")
        return telemetry

    return telemetry


if __name__ == "__main__":
    is_dry = "--dry-run" in sys.argv
    res = check_and_run_one(dry_run=is_dry)
    if is_dry:
        print(json.dumps(res, indent=2))
