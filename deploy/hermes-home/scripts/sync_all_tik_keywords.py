"""
Script & Cron Watchdog: Tự động quét và đồng bộ Niche, Keyword, Hashtag Pool từ state.db
sang TOÀN BỘ CÁC FILE TIK (Tik1 -> Tik8) cho CẢ 2 CLUSTER (Kibe & Admin).

Mục đích:
- Tự động phát hiện khi một folder video gốc (1..640) được gán/đổi niche mới trong state.db.
- Nếu nick bị xóa, đổi nội dung hoặc tạo bộ video mới, script sẽ tự động cập nhật lại đúng
  Keyword Video và Hashtag Pool vào file Tik tương ứng của cả Kibe và Admin.
- Đảm bảo tính nhất quán tuyệt đối giữa:
  `state.db` -> `niches_pool.txt` -> `Hashtag theo Folder` -> `TaiKhoan` (cột Keyword & Hashtag).
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import logging
import os
import re
import sqlite3
import sys
from pathlib import Path
from typing import Any

logger = logging.getLogger("sync_all_tik_keywords")

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    core_path = Path("D:/Taadaa/automation-core/src")
    if core_path.exists() and str(core_path) not in sys.path:
        sys.path.insert(0, str(core_path))
    from automation_core.workbook import atomic_workbook_update
except ImportError:
    atomic_workbook_update = None

STATE_DB = Path(r"D:\CodexRuntime\tiktok-video\state.db")
NICHES_POOL_FILE = Path(r"D:\Taadaa\Tiktok-video\data\niches_pool.txt")
DEFAULT_TELEMETRY_LOG = Path("D:/Taadaa/logs/sync_keywords.jsonl")

FARM_DIRS = [
    ("Kibe", Path(r"D:\OneDrive\TaadaaData\kibe")),
    ("Admin", Path(r"D:\OneDrive\TaadaaData\admin")),
]

# 8 file Tik tương ứng 8 slot
TIK_FILES = [
    ("Tik1.xlsx", 1, 0),     # slot 1, base source: (1-1)*80 = 0 -> 1..80
    ("Tik2.xlsx", 2, 80),    # slot 2, base source: (2-1)*80 = 80 -> 81..160
    ("tik3.xlsx", 3, 160),   # slot 3, base source: (3-1)*80 = 160 -> 161..240
    ("Tik4.xlsx", 4, 240),   # slot 4, base source: (4-1)*80 = 240 -> 241..320
    ("Tik5.xlsx", 5, 320),   # slot 5, base source: (5-1)*80 = 320 -> 321..400
    ("Tik6.xlsx", 6, 400),   # slot 6, base source: (6-1)*80 = 400 -> 401..480
    ("Tik7.xlsx", 7, 480),   # slot 7, base source: (7-1)*80 = 480 -> 481..560
    ("Tik8.xlsx", 8, 560),   # slot 8, base source: (8-1)*80 = 560 -> 561..640
]

# Machine range and offset definitions across dual clusters
KIBE_MAX_MACHINE = 80
ADMIN_MIN_MACHINE = 201
ADMIN_MAX_MACHINE = 280
ADMIN_MACHINE_OFFSET = 200


def resolve_source_folder(slot_num: int, machine_val: Any) -> int | None:
    """
    Calculate source folder F_src from slot (1..8) and machine number.
    - Kibe: machines 1..80 -> m_idx = machine (1..80)
    - Admin: machines 201..280 -> m_idx = machine - 200 (1..80)
    Formula: F_src = (slot_num - 1) * 80 + m_idx (range: 1..640)
    """
    if machine_val is None or not (1 <= slot_num <= 8):
        return None
    if isinstance(machine_val, float) and not machine_val.is_integer():
        return None
    if isinstance(machine_val, str) and not machine_val.strip().isdigit():
        return None
    try:
        m_int = int(machine_val)
    except (ValueError, TypeError):
        return None

    if 1 <= m_int <= KIBE_MAX_MACHINE:
        m_idx = m_int
    elif ADMIN_MIN_MACHINE <= m_int <= ADMIN_MAX_MACHINE:
        m_idx = m_int - ADMIN_MACHINE_OFFSET
    else:
        return None
    return (slot_num - 1) * 80 + m_idx


def load_niche_definitions() -> dict[str, str]:
    """Đọc niches_pool.txt để lấy slug -> label tiếng Việt."""
    niche_labels: dict[str, str] = {}
    if not NICHES_POOL_FILE.exists():
        return niche_labels

    with open(NICHES_POOL_FILE, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) >= 3:
                slug = parts[1].strip()
                label = parts[2].strip()
                niche_labels[slug] = label
    return niche_labels


def load_existing_hashtag_pools() -> dict[str, str]:
    """Thu thập toàn bộ mapping mẫu (Keyword/Label -> Hashtag Pool) từ các file Tik đã chuẩn hóa."""
    keyword_to_pool: dict[str, str] = {}
    if openpyxl is None:
        return keyword_to_pool

    kibe_dir = Path(r"D:\OneDrive\TaadaaData\kibe")
    if not kibe_dir.exists():
        return keyword_to_pool

    for filename, _, _ in TIK_FILES[:6]:
        p = kibe_dir / filename
        if not p.exists():
            continue
        try:
            wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
            if "Hashtag theo Folder" in wb.sheetnames:
                ws = wb["Hashtag theo Folder"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 4 and row[2] and row[3]:
                        kw = str(row[2]).strip().lower()
                        tags = str(row[3]).strip()
                        if kw and tags and kw not in keyword_to_pool:
                            keyword_to_pool[kw] = tags
            wb.close()
        except (OSError, KeyError, ValueError) as err:
            logger.warning(f"Could not load hashtag pools from {p}: {err}")
        except Exception as err:
            logger.warning(f"Unexpected error loading hashtag pools from {p}: {err}")

    return keyword_to_pool


def generate_hashtag_pool(slug: str, label: str, existing_pools: dict[str, str]) -> str:
    """Tạo hoặc lấy pool hashtag chuẩn cho 1 niche."""
    kw_key = label.strip().lower()
    if kw_key in existing_pools:
        return existing_pools[kw_key]

    clean_slug = re.sub(r"[^a-zA-Z0-9_]", "", slug).lower()
    pool = (
        f"#{clean_slug} #{clean_slug}vietnam #{clean_slug}moingay "
        f"#tiktokvietnam #xuhuong #fyp #videohay"
    )
    return pool


def get_state_db_niches() -> dict[int, tuple[str, int, str]]:
    """Đọc từ state.db: folder_num -> (niche, video_count, status)."""
    db_map: dict[int, tuple[str, int, str]] = {}
    if not STATE_DB.exists():
        return db_map

    try:
        conn = sqlite3.connect(str(STATE_DB))
        cursor = conn.cursor()
        query = (
            "SELECT folder_num, niche, video_count, status "
            "FROM folders "
            "WHERE niche IS NOT NULL AND niche != ''"
        )
        for r in cursor.execute(query).fetchall():
            f_num = int(r[0])
            niche = str(r[1]).strip()
            v_count = int(r[2]) if r[2] is not None else 0
            status = str(r[3] or "pending").strip()
            db_map[f_num] = (niche, v_count, status)
        conn.close()
    except Exception as e:
        sys.stderr.write(f"Error querying state.db: {e}\n")

    return db_map


def sync_tik_file(
    folder_path: Path,
    filename: str,
    slot_num: int,
    base_source: int,
    db_niches: dict[int, tuple[str, int, str]],
    niche_labels: dict[str, str],
    existing_pools: dict[str, str],
    stats: dict[str, Any] | None = None,
) -> int:
    """Đồng bộ 1 file Tik."""
    p = folder_path / filename
    if not p.exists() and folder_path.exists():
        target_lower = filename.lower()
        for f in folder_path.iterdir():
            if f.name.lower() == target_lower:
                p = f
                break
    if not p.exists() or openpyxl is None:
        return 0

    if stats is not None:
        stats["workbooks_checked"] = stats.get("workbooks_checked", 0) + 1

    def _update_workbook(target_path: Path) -> int:
        try:
            wb = openpyxl.load_workbook(str(target_path))
        except (OSError, KeyError, ValueError) as e:
            logger.error(f"Cannot open {filename}: {e}")
            sys.stderr.write(f"Cannot open {filename}: {e}\n")
            if stats is not None:
                stats["error_count"] = stats.get("error_count", 0) + 1
            return 0
        except Exception as e:
            logger.error(f"Unexpected error opening {filename}: {e}")
            sys.stderr.write(f"Cannot open {filename}: {e}\n")
            if stats is not None:
                stats["error_count"] = stats.get("error_count", 0) + 1
            return 0

        updates = 0
        try:
            ws_tai_khoan = wb["TaiKhoan"] if "TaiKhoan" in wb.sheetnames else wb.active
            headers = [str(c or "").strip().lower() for c in next(ws_tai_khoan.iter_rows(max_row=1, values_only=True))]

            may_col = headers.index("máy") if "máy" in headers else -1
            vg_col = headers.index("video gốc") if "video gốc" in headers else 4
            kw_col = headers.index("keyword video") if "keyword video" in headers else 5
            ht_col = headers.index("hashtag pool") if "hashtag pool" in headers else 6

            # 0. Build cache from Hashtag theo Folder if available
            ht_map = {}
            if "Hashtag theo Folder" in wb.sheetnames:
                ws_ht_preload = wb["Hashtag theo Folder"]
                for r in ws_ht_preload.iter_rows(min_row=2, values_only=True):
                    if r[0] is not None and len(r) >= 4 and r[2] and r[3]:
                        try:
                            ht_map[int(r[0])] = (str(r[2]).strip(), str(r[3]).strip())
                        except (ValueError, TypeError) as err:
                            logger.warning(f"Invalid folder id row in 'Hashtag theo Folder' {filename}: {err}")

            # 1. Update sheet TaiKhoan
            for row in range(2, ws_tai_khoan.max_row + 1):
                m_val = ws_tai_khoan.cell(row, may_col + 1).value if may_col != -1 else None
                f_src = resolve_source_folder(slot_num, m_val)

                if f_src is None:
                    if m_val is not None:
                        logger.warning(
                            f"[{filename}] Row {row}: Invalid machine value '{m_val}' for slot {slot_num}"
                        )
                    vg_val = ws_tai_khoan.cell(row, vg_col + 1).value
                    if vg_val is None:
                        if stats is not None:
                            stats["skipped_rows"] = stats.get("skipped_rows", 0) + 1
                        continue
                    try:
                        f_src = int(vg_val)
                    except (ValueError, TypeError) as err:
                        logger.warning(
                            f"[{filename}] Row {row}: Invalid 'video gốc' value '{vg_val}': {err}"
                        )
                        if stats is not None:
                            stats["skipped_rows"] = stats.get("skipped_rows", 0) + 1
                        continue

                curr_vg = ws_tai_khoan.cell(row, vg_col + 1).value
                if curr_vg != f_src:
                    ws_tai_khoan.cell(row, vg_col + 1, f_src)
                    updates += 1

                label = None
                pool = None
                if f_src in db_niches:
                    niche_slug, _, _ = db_niches[f_src]
                    label = niche_labels.get(niche_slug, niche_slug.capitalize())
                    pool = generate_hashtag_pool(niche_slug, label, existing_pools)
                elif f_src in ht_map:
                    label, pool = ht_map[f_src]

                if label and pool:
                    curr_kw = ws_tai_khoan.cell(row, kw_col + 1).value
                    curr_ht = ws_tai_khoan.cell(row, ht_col + 1).value

                    if curr_kw != label or curr_ht != pool:
                        ws_tai_khoan.cell(row, kw_col + 1, label)
                        ws_tai_khoan.cell(row, ht_col + 1, pool)
                        updates += 1

            # 2. Update sheet Hashtag theo Folder nếu có
            if "Hashtag theo Folder" in wb.sheetnames:
                ws_ht = wb["Hashtag theo Folder"]
                for row in range(2, ws_ht.max_row + 1):
                    src_val = ws_ht.cell(row, 1).value
                    if src_val is None:
                        continue
                    try:
                        f_src = int(src_val)
                    except (ValueError, TypeError):
                        continue

                    if f_src in db_niches:
                        niche_slug, v_count, status = db_niches[f_src]
                        label = niche_labels.get(niche_slug, niche_slug.capitalize())
                        pool = generate_hashtag_pool(niche_slug, label, existing_pools)

                        c_label = ws_ht.cell(row, 3).value
                        c_pool = ws_ht.cell(row, 4).value
                        c_count = ws_ht.cell(row, 5).value
                        c_status = ws_ht.cell(row, 6).value

                        if (
                            c_label != label
                            or c_pool != pool
                            or c_count != v_count
                            or c_status != status
                        ):
                            ws_ht.cell(row, 3, label)
                            ws_ht.cell(row, 4, pool)
                            ws_ht.cell(row, 5, v_count)
                            ws_ht.cell(row, 6, status)
                            ws_ht.cell(row, 7, "OK" if "complete" in status.lower() else status.upper())
                            updates += 1

            if updates > 0:
                wb.save(str(target_path))
        finally:
            wb.close()
        return updates

    try:
        if atomic_workbook_update is not None:
            return atomic_workbook_update(p, _update_workbook, backup=True)
        return _update_workbook(p)
    except Exception as e:
        sys.stderr.write(f"Failed to update {filename}: {e}\n")
        return 0


def build_sync_telemetry(
    clusters: list[str],
    total_updated: int,
    details: list[str],
    stats: dict[str, Any],
) -> dict[str, Any]:
    """Tạo payload telemetry có cấu trúc chuẩn JSON."""
    return {
        "event": "SYNC_KEYWORDS_SUMMARY",
        "correlation_id": f"sync_{int(datetime.now(timezone.utc).timestamp())}",
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "clusters": clusters,
        "total_updated": total_updated,
        "workbooks_checked": stats.get("workbooks_checked", 0),
        "skipped_rows": stats.get("skipped_rows", 0),
        "error_count": stats.get("error_count", 0),
        "details": details,
        "status": "SUCCESS" if stats.get("error_count", 0) == 0 else "WARNING",
    }


def emit_sync_telemetry(telemetry: dict[str, Any], log_path: Path | str | None = None) -> Path | None:
    """Ghi append 1 dòng JSON telemetry vào file jsonl; trả về đường dẫn đã ghi hoặc None."""
    if log_path is None:
        if not DEFAULT_TELEMETRY_LOG.parent.exists():
            return None
        log_path = DEFAULT_TELEMETRY_LOG
    target = Path(log_path)
    try:
        with open(target, "a", encoding="utf-8") as f:
            f.write(json.dumps(telemetry, ensure_ascii=False) + "\n")
    except OSError as err:
        logger.error(f"Cannot write sync telemetry to {target}: {err}")
        sys.stderr.write(f"Cannot write sync telemetry to {target}: {err}\n")
        return None
    return target


def main() -> int:
    parser = argparse.ArgumentParser(description="Auto sync keywords & hashtags from state.db to all Tik workbooks.")
    parser.add_argument("--silent", action="store_true", help="Do not print if no updates")
    args = parser.parse_args()

    niche_labels = load_niche_definitions()
    existing_pools = load_existing_hashtag_pools()
    db_niches = get_state_db_niches()

    total_updated = 0
    details = []
    stats: dict[str, Any] = {
        "workbooks_checked": 0,
        "skipped_rows": 0,
        "error_count": 0,
    }

    for cluster_name, cluster_dir in FARM_DIRS:
        if not cluster_dir.exists():
            continue
        for filename, slot_num, base_source in TIK_FILES:
            n = sync_tik_file(
                cluster_dir,
                filename,
                slot_num,
                base_source,
                db_niches,
                niche_labels,
                existing_pools,
                stats=stats,
            )
            if n > 0:
                details.append(f"{cluster_name}/{filename}: {n} updates")
                total_updated += n

    telemetry = build_sync_telemetry([c[0] for c in FARM_DIRS], total_updated, details, stats)
    logger.info(json.dumps(telemetry))
    emit_sync_telemetry(telemetry)

    if total_updated > 0:
        print(f"[SYNC_KEYWORDS] Synced {total_updated} changes across workbooks: {', '.join(details)}")
    elif not args.silent:
        print("[SYNC_KEYWORDS] All Tik workbooks (Kibe & Admin) are already up to date with state.db.")

    return 0


if __name__ == "__main__":
    sys.exit(main())
